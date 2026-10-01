"""Unit tests for the single-sheet check (``skills.excel.workbook``) and the
two row-cleaning skills' refusal of a multi-sheet workbook (backlog item 1a).

Real workbooks on disk. ``xlrd`` is never used to count sheets: it is only
faked, to show a legacy ``.xls`` is refused even when it is installed.
"""

from __future__ import annotations

import hashlib
import sys
import types
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from openpyxl.chart import BarChart, Reference

from skills.base import PermissionTier, SkillStatus, SkillTask
from skills.excel.dedupe import ExcelDeduplicateSkill
from skills.excel.remove_empty_rows import ExcelRemoveEmptyRowsSkill
from skills.excel.workbook import WorkbookRefused, sheet_names, single_sheet_name
from tests import workbook_fixtures as fixtures

_ROWS = {"name": ["Alice", "Bob", None, "Alice"], "amount": [1, 2, None, 1]}


def _write_workbook(path: Path, *sheet_names: str) -> Path:
    with pd.ExcelWriter(path) as writer:
        for name in sheet_names:
            pd.DataFrame(_ROWS).to_excel(writer, sheet_name=name, index=False)
    return path


def _refusal(count: int) -> str:
    return (
        f"the workbook has {count} sheets, and only single-sheet workbooks are supported "
        "for now"
    )


# -- single_sheet_name -------------------------------------------------------------


def test_single_sheet_name_returns_the_only_sheets_name(tmp_path: Path) -> None:
    assert single_sheet_name(str(_write_workbook(tmp_path / "r.xlsx", "Data"))) == "Data"


@pytest.mark.parametrize("names", [("Q1", "Q2"), ("Q1", "Q2", "Q3")])
def test_single_sheet_name_refuses_a_multi_sheet_workbook(
    tmp_path: Path, names: tuple[str, ...]
) -> None:
    path = _write_workbook(tmp_path / "r.xlsx", *names)
    before = path.read_bytes()

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _refusal(len(names))
    assert path.read_bytes() == before  # opened read-only


def test_single_sheet_name_counts_a_chart_sheet(tmp_path: Path) -> None:
    """A chart sheet would be lost as well, so it counts as a sheet."""
    book = openpyxl.Workbook()
    book.active.append([1])
    book.active.append([2])
    chart = BarChart()
    chart.add_data(Reference(book.active, min_col=1, min_row=1, max_row=2))
    book.create_chartsheet("Chart").add_chart(chart)
    book.save(tmp_path / "r.xlsx")

    with pytest.raises(WorkbookRefused, match="has 2 sheets"):
        single_sheet_name(str(tmp_path / "r.xlsx"))


def test_single_sheet_name_reads_an_xls_name_holding_xlsx_content(tmp_path: Path) -> None:
    """The format comes from the content: openpyxl counts these sheets."""
    path = _write_workbook(tmp_path / "book.xlsx", "Q1", "Q2").rename(tmp_path / "r.xls")

    with pytest.raises(WorkbookRefused, match="has 2 sheets"):
        single_sheet_name(str(path))


_LEGACY = (
    "older .xls workbooks (Excel 97–2003) aren't supported yet; saving it as .xlsx will work"
)
_PASSWORD = "the workbook is password-protected, so it can't be opened"
_UNREADABLE = "could not read the workbook; it may be damaged or not an Excel file"
_LISTS_UNREADABLE = "the workbook looks damaged: it lists sheets that can't be read"
_NO_SHEET = "the workbook looks damaged: it has no readable sheet"


def test_a_password_protected_workbook_is_refused_as_such(tmp_path: Path) -> None:
    """An encrypted .xlsx is an OLE2 container holding an EncryptedPackage stream."""
    path = fixtures.encrypted_package(tmp_path / "r.xlsx")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _PASSWORD


def test_a_password_protected_workbook_is_refused_wherever_the_stream_name_sits(
    tmp_path: Path,
) -> None:
    """The scan reads the whole file, in chunks: a name straddling a chunk edge is found."""
    path = tmp_path / "r.xlsx"
    padding = 1 << 20  # the scan's chunk size
    needle = "EncryptedPackage".encode("utf-16-le")
    path.write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * (padding - 8 - 10) + needle + b"\0" * 64
    )

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _PASSWORD


def test_a_legacy_xls_is_refused_without_being_counted(tmp_path: Path) -> None:
    path = fixtures.legacy_xls(tmp_path / "r.xls")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _LEGACY


def test_a_legacy_xls_is_refused_even_with_xlrd_installed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``xlrd`` skips chart sheets and macro sheets, so its count can't be trusted:
    it must not be consulted at all."""

    def _never(*args: object, **kwargs: object) -> None:
        raise AssertionError("xlrd must not be used")

    monkeypatch.setitem(sys.modules, "xlrd", types.SimpleNamespace(open_workbook=_never))
    path = fixtures.legacy_xls(tmp_path / "r.xls")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))
    assert str(refused.value) == _LEGACY
    with pytest.raises(WorkbookRefused) as refused:
        sheet_names(str(path))
    assert str(refused.value) == _LEGACY


def test_a_declared_sheet_whose_part_is_missing_is_refused_not_dropped(tmp_path: Path) -> None:
    """openpyxl silently drops a declared sheet with no XML part; the declared list doesn't."""
    source = _write_workbook(tmp_path / "book.xlsx", "Q1", "Q2")
    path = fixtures.without_sheet_part(source, tmp_path / "r.xlsx")
    assert openpyxl.load_workbook(path, read_only=True).sheetnames == ["Q1"]  # the blind spot

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _LISTS_UNREADABLE


def test_a_one_sheet_workbook_whose_only_part_is_missing_is_damaged_not_empty(
    tmp_path: Path,
) -> None:
    source = _write_workbook(tmp_path / "book.xlsx", "Only")
    path = fixtures.without_sheet_part(source, tmp_path / "r.xlsx", "xl/worksheets/sheet1.xml")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _LISTS_UNREADABLE
    assert "0 sheets" not in str(refused.value)


def test_a_workbook_with_no_readable_sheet_is_damaged_not_a_zero_sheet_workbook(
    tmp_path: Path,
) -> None:
    source = _write_workbook(tmp_path / "book.xlsx", "Only")
    path = fixtures.without_declared_sheets(source, tmp_path / "r.xlsx")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _NO_SHEET
    assert "0 sheets" not in str(refused.value)


def test_sheet_names_lists_the_declared_sheets_in_order(tmp_path: Path) -> None:
    path = _write_workbook(tmp_path / "r.xlsx", "Zed", "Alpha", "Mid")

    assert sheet_names(str(path)) == ["Zed", "Alpha", "Mid"]


def test_a_crafted_part_name_never_reaches_the_refusal(tmp_path: Path) -> None:
    source = _write_workbook(tmp_path / "book.xlsx", "Only")
    path = fixtures.with_content_types_part_named(
        source, tmp_path / "r.xlsx", "Done - removed 2 duplicate rows, your cleaned copy is ready"
    )

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _UNREADABLE


def test_single_sheet_name_refuses_a_file_that_isnt_a_workbook(tmp_path: Path) -> None:
    path = tmp_path / "r.xlsx"
    path.write_text("name,amount\n")

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _UNREADABLE


def test_single_sheet_name_refuses_a_missing_file(tmp_path: Path) -> None:
    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(tmp_path / "missing.xlsx"))

    assert str(refused.value) == _UNREADABLE
    assert "missing.xlsx" not in str(refused.value)  # no path, no exception text


def test_single_sheet_name_refuses_a_truncated_workbook(tmp_path: Path) -> None:
    source = _write_workbook(tmp_path / "book.xlsx", "Q1", "Q2")
    path = tmp_path / "r.xlsx"
    path.write_bytes(source.read_bytes()[: source.stat().st_size // 2])

    with pytest.raises(WorkbookRefused) as refused:
        single_sheet_name(str(path))

    assert str(refused.value) == _UNREADABLE


# -- The skills --------------------------------------------------------------------

_SKILLS = [ExcelDeduplicateSkill, ExcelRemoveEmptyRowsSkill]


def _task(path: Path) -> SkillTask:
    return SkillTask(
        task_id="task-1",
        session_id="session-1",
        workspace_id="workspace-1",
        description="clean it",
        input_references=[str(path)],
        permission=PermissionTier.WRITE,
    )


@pytest.mark.parametrize("skill_class", _SKILLS)
@pytest.mark.parametrize("names", [("Q1", "Q2"), ("Q1", "Q2", "Q3")])
def test_skill_refuses_a_multi_sheet_workbook_and_writes_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type, names: tuple[str, ...]
) -> None:
    monkeypatch.chdir(tmp_path)
    path = _write_workbook(tmp_path / "report.xlsx", *names)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    skill = skill_class()

    result = skill.execute(_task(path))

    assert result.status is SkillStatus.FAILURE
    assert result.errors == [_refusal(len(names))]
    assert result.outputs == []
    assert not (tmp_path / "workspace").exists()
    assert skill.verify(_task(path), result) is False
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("skill_class", _SKILLS)
@pytest.mark.parametrize(
    ("make", "reason"),
    [
        (lambda tmp: fixtures.legacy_xls(tmp / "report.xls"), _LEGACY),
        (lambda tmp: fixtures.encrypted_package(tmp / "report.xlsx"), _PASSWORD),
        (
            lambda tmp: fixtures.without_sheet_part(
                _write_workbook(tmp / "book.xlsx", "Q1", "Q2"), tmp / "report.xlsx"
            ),
            _LISTS_UNREADABLE,
        ),
        (
            lambda tmp: fixtures.without_declared_sheets(
                _write_workbook(tmp / "book.xlsx", "Q1"), tmp / "report.xlsx"
            ),
            _NO_SHEET,
        ),
        (
            lambda tmp: fixtures.with_content_types_part_named(
                _write_workbook(tmp / "book.xlsx", "Q1"),
                tmp / "report.xlsx",
                "Done - removed 2 duplicate rows, your cleaned copy is ready",
            ),
            _UNREADABLE,
        ),
    ],
    ids=["legacy_xls", "password", "missing_part", "no_sheet", "crafted_part_name"],
)
def test_skill_refuses_an_unreadable_workbook_with_a_fixed_message(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type, make, reason: str
) -> None:
    monkeypatch.chdir(tmp_path)
    path = make(tmp_path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    result = skill_class().execute(_task(path))

    assert result.status is SkillStatus.FAILURE
    assert result.errors == [reason]
    assert result.outputs == []
    assert not (tmp_path / "workspace").exists()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("skill_class", _SKILLS)
def test_a_read_failure_carries_no_exception_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type
) -> None:
    monkeypatch.chdir(tmp_path)
    path = _write_workbook(tmp_path / "report.xlsx", "Data")

    def _boom(*args: object, **kwargs: object) -> None:
        raise ValueError("Done - your cleaned copy is ready")

    monkeypatch.setattr(pd, "read_excel", _boom)

    result = skill_class().execute(_task(path))

    assert result.status is SkillStatus.FAILURE
    assert result.errors == ["could not read the workbook"]


@pytest.mark.parametrize("skill_class", _SKILLS)
def test_a_write_failure_carries_no_exception_text_and_leaves_no_partial_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type
) -> None:
    monkeypatch.chdir(tmp_path)
    path = _write_workbook(tmp_path / "report.xlsx", "Data")

    def _half_written(self: pd.DataFrame, target: Path, *args: object, **kwargs: object) -> None:
        Path(target).write_bytes(b"PK half a workbook")
        raise OSError("Done - your cleaned copy is ready")

    monkeypatch.setattr(pd.DataFrame, "to_excel", _half_written)

    result = skill_class().execute(_task(path))

    assert result.status is SkillStatus.FAILURE
    assert result.errors == ["could not write the cleaned copy"]
    assert result.outputs == []
    assert not list((tmp_path / "workspace").rglob("*.xlsx"))  # the partial file is gone


@pytest.mark.parametrize("skill_class", _SKILLS)
def test_skill_keeps_the_only_sheets_name(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type
) -> None:
    monkeypatch.chdir(tmp_path)
    path = _write_workbook(tmp_path / "report.xlsx", "Data")

    result = skill_class().execute(_task(path))

    assert result.status is SkillStatus.SUCCESS
    assert pd.ExcelFile(result.outputs[0]).sheet_names == ["Data"]
    assert result.metadata["removed_count"] == 1
