"""The plain-table check (``docs/design/values-only-copy.md``, Section 1) and the formula count.

``plain_table_problem`` / ``plain_table_sheet_name`` read the file with openpyxl, never a
skill's DataFrame; the formula counts are checked here against a raw read of the sheet's XML.
"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pytest

from skills.base import SkillStatus, SkillTask
from skills.excel import workbook
from skills.excel.dedupe import ExcelDeduplicateSkill
from skills.excel.remove_empty_rows import ExcelRemoveEmptyRowsSkill
from skills.excel.workbook import (
    FormulaCounts,
    WorkbookRefused,
    formula_cell_counts,
    plain_table_problem,
    plain_table_sheet_name,
)
from tests import workbook_fixtures as fixtures

_SUFFIX = ", and only plain tables (one header row at the top, then the data) are supported for now"
_MERGED = "the sheet has merged cells" + _SUFFIX
_HEADER_NOT_FIRST = "the sheet's header isn't on the first row" + _SUFFIX
_NO_HEADER = "a column in the sheet has no header" + _SUFFIX
_DUPLICATE = "two columns in the sheet have the same header" + _SUFFIX
_READS_AS_EMPTY = (
    "a cell holds an error value, or text such as NA or N/A, that would be read as empty, "
    "and keeping such cells isn't supported for now"
)


def _book(path: Path, rows: list[list[object]], *, title: str = "Sheet1") -> Path:
    """A workbook with one sheet whose cells are ``rows`` (``None`` = nothing)."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=row_index, column=column_index, value=value)
    book.save(path)
    return path


def _independent_formula_counts(path: Path) -> tuple[int, int]:
    """Formula cells, and those with no saved value, read straight from every sheet's XML."""
    formulas = without_value = 0
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/sheet"):
                continue
            for cell in ElementTree.fromstring(archive.read(name)).iter():
                if not cell.tag.endswith("}c"):
                    continue
                if any(child.tag.endswith("}f") for child in cell):
                    formulas += 1
                    value = next((c for c in cell if c.tag.endswith("}v")), None)
                    without_value += value is None or not (value.text or "")
    return formulas, without_value


# -- plain tables are accepted -----------------------------------------------------------


@pytest.mark.parametrize(
    "rows",
    [
        [["name", "amount"], ["Alice", 1], ["Bob", 2]],
        [["name", "amount"]],  # a header and no data
        [["name"], ["Alice"], [None], ["Bob"]],  # a one-column table with an empty row
        [["name", "amount"], [None, 5], ["Bob", None]],  # a gap in the data
        [[2019, 2020], [1, 2]],  # numeric headers
        [["name", " "], ["Alice", 1]],  # a blank-looking header is still a header
        [["name", "amount", "double"], ["A", 1, "=B2*2"]],  # formulas are not a layout problem
        [],  # an empty sheet: nothing to change
    ],
)
def test_a_plain_table_has_no_problem(tmp_path: Path, rows: list[list[object]]) -> None:
    path = _book(tmp_path / "r.xlsx", rows)
    assert plain_table_problem(str(path)) is None
    assert plain_table_sheet_name(str(path)) == "Sheet1"


def test_a_formatted_sheet_with_formulas_is_still_a_plain_table(tmp_path: Path) -> None:
    path = fixtures.formatted_formula_sheet(tmp_path / "r.xlsx", saved=True)
    assert plain_table_problem(str(path)) is None


# -- what isn't a plain table ------------------------------------------------------------


def test_merged_cells_are_refused(tmp_path: Path) -> None:
    path = fixtures.merged_cell_sheet(tmp_path / "r.xlsx")
    assert plain_table_problem(str(path)) == _MERGED
    with pytest.raises(WorkbookRefused) as refused:
        plain_table_sheet_name(str(path))
    assert str(refused.value) == _MERGED


@pytest.mark.parametrize("merged_range", ["A1:B1", "A3:B3", "A2:A3", "E5:G9"])
def test_a_merged_range_anywhere_is_refused(tmp_path: Path, merged_range: str) -> None:
    path = _book(tmp_path / "r.xlsx", [["name", "amount"], ["Alice", 1], ["Bob", 2]])
    book = openpyxl.load_workbook(path)
    book.active.merge_cells(merged_range)
    book.save(path)
    assert plain_table_problem(str(path)) == _MERGED


def test_a_merge_cell_after_a_large_sheet_is_found_across_read_chunks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The scan reads the sheet's XML a chunk at a time; a tag split between chunks is still seen.
    monkeypatch.setattr(workbook, "_SCAN_CHUNK_BYTES", 7)
    path = fixtures.merged_cell_sheet(tmp_path / "r.xlsx")
    assert plain_table_problem(str(path)) == _MERGED


def test_a_header_below_a_blank_first_row_is_refused(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [[], [], ["name", "amount"], ["Alice", 1]])
    assert plain_table_problem(str(path)) == _HEADER_NOT_FIRST


def test_a_table_that_starts_below_row_one_is_refused(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [[], ["name", "amount"], ["Alice", 1]])
    assert plain_table_problem(str(path)) == _HEADER_NOT_FIRST


def test_a_header_on_row_three_under_a_title_is_refused(tmp_path: Path) -> None:
    path = fixtures.header_on_row_three_sheet(tmp_path / "r.xlsx")
    assert plain_table_problem(str(path)) == _NO_HEADER


@pytest.mark.parametrize(
    "rows",
    [
        [["name", None, "amount"], ["Alice", 1, 2]],  # a gap in the header row
        [[None, "amount"], ["Alice", 1]],  # the first column has no header
        [["name", "amount"], ["Alice", 1, "extra"]],  # data to the right of the last header
        [["name", ""], ["Alice", 1]],  # an empty-string header is no header
        [["Title"], [None, "x", "y"], ["a", 1, 2]],
    ],
)
def test_a_column_without_a_header_is_refused(tmp_path: Path, rows: list[list[object]]) -> None:
    path = _book(tmp_path / "r.xlsx", rows)
    assert plain_table_problem(str(path)) == _NO_HEADER


def test_an_empty_string_header_written_as_a_string_cell_is_refused(tmp_path: Path) -> None:
    # A shared/inline string cell holding "" is a cell with a value of "", which pandas reads
    # as missing; openpyxl itself drops such a cell on write, so the XML is crafted.
    path = _book(tmp_path / "r.xlsx", [["name", "marker"], ["Alice", 1]])
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    assert "<t>marker</t>" in xml
    crafted = fixtures.rewrite(
        path, tmp_path / "e.xlsx", {"xl/worksheets/sheet1.xml": xml.replace("marker", "").encode()}
    )
    assert plain_table_problem(str(crafted)) == _NO_HEADER


def test_two_columns_with_the_same_header_are_refused(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [["Total", "amount", "Total"], [1, 2, 3]])
    assert plain_table_problem(str(path)) == _DUPLICATE


def test_the_refusal_text_never_carries_a_name_or_value_from_the_file(tmp_path: Path) -> None:
    rows = [["SecretHeader", "SecretHeader"], ["SecretValue", 1]]
    path = _book(tmp_path / "r.xlsx", rows, title="SecretSheet")
    message = plain_table_problem(str(path))
    assert message == _DUPLICATE
    assert "Secret" not in message


def test_a_multi_sheet_workbook_is_refused_for_its_sheets_first(tmp_path: Path) -> None:
    path = fixtures.merged_cell_sheet(tmp_path / "r.xlsx")
    book = openpyxl.load_workbook(path)
    book.create_sheet("Other")
    book.save(path)
    with pytest.raises(WorkbookRefused, match="has 2 sheets"):
        plain_table_sheet_name(str(path))


def test_plain_table_problem_looks_at_every_sheet(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [["name"], ["Alice"]])
    book = openpyxl.load_workbook(path)
    second = book.create_sheet("Other")
    second.append(["a", "a"])
    book.save(path)
    assert plain_table_problem(str(path)) == _DUPLICATE


def test_an_unreadable_file_is_refused_with_fixed_text(tmp_path: Path) -> None:
    path = tmp_path / "r.xlsx"
    path.write_bytes(b"not a workbook")
    for check in (plain_table_problem, plain_table_sheet_name, formula_cell_counts):
        with pytest.raises(WorkbookRefused) as refused:
            check(str(path))
        assert "not a workbook" not in str(refused.value)


# -- cells pandas reads as missing although they aren't empty (review round 1, R1-1 and R1-2) --


def test_the_na_text_set_is_pandas_own() -> None:
    from pandas._libs.parsers import STR_NA_VALUES

    assert workbook._PANDAS_NA_TEXT == frozenset(STR_NA_VALUES) - {""}


@pytest.mark.parametrize("text", sorted(workbook._PANDAS_NA_TEXT))
def test_text_pandas_reads_as_missing_is_refused_in_a_data_cell(tmp_path: Path, text: str) -> None:
    path = _book(tmp_path / "r.xlsx", [["country", "note"], ["US", text]])
    assert plain_table_problem(str(path)) == _READS_AS_EMPTY


def test_text_pandas_reads_as_missing_is_allowed_in_a_header(tmp_path: Path) -> None:
    # pandas keeps a header as written, so it comes back unchanged.
    path = _book(tmp_path / "r.xlsx", [["NA", "N/A"], ["US", "x"]])
    assert plain_table_problem(str(path)) is None


def test_ordinary_text_close_to_na_is_not_refused(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [["country", "note"], ["Namibia", "n/a."], ["na", "NAN"]])
    assert plain_table_problem(str(path)) is None


def test_a_typed_error_cell_is_refused(tmp_path: Path) -> None:
    path = _book(tmp_path / "r.xlsx", [["name", "amount"], ["Alice", "#DIV/0!"]])
    assert plain_table_problem(str(path)) == _READS_AS_EMPTY


def test_a_formula_with_a_saved_error_result_is_refused(tmp_path: Path) -> None:
    path = fixtures.error_result_sheet(tmp_path / "r.xlsx")
    assert plain_table_problem(str(path)) == _READS_AS_EMPTY


def test_the_na_sheet_is_refused_by_the_skills(tmp_path: Path) -> None:
    path = fixtures.na_text_sheet(tmp_path / "uploads" / "report.xlsx")
    task = SkillTask(
        task_id="t",
        session_id="s",
        workspace_id="w",
        description="clean it",
        input_references=[str(path)],
        output_dir=str(tmp_path / "out"),
    )
    for skill_class in (ExcelDeduplicateSkill, ExcelRemoveEmptyRowsSkill):
        result = skill_class().execute(task)
        assert result.status is SkillStatus.FAILURE
        assert result.errors == [_READS_AS_EMPTY]
    assert not (tmp_path / "out").exists()


# -- a wrong <dimension> doesn't hide cells (review round 1, R1-3) -----------------------


@pytest.mark.parametrize("ref", ["A1", "A1:A2", "A1:B2"])
def test_a_wrong_dimension_doesnt_hide_a_header_gap(tmp_path: Path, ref: str) -> None:
    path = fixtures.gap_header_wrong_dimension_sheet(tmp_path / "r.xlsx")
    wrong = fixtures.with_dimension(path, tmp_path / "w.xlsx", ref)
    assert plain_table_problem(str(wrong)) == _NO_HEADER


def test_a_wrong_dimension_doesnt_hide_formulas_from_the_count(tmp_path: Path) -> None:
    plain = fixtures.formatted_formula_sheet(tmp_path / "r.xlsx", saved=False)
    wrong = fixtures.with_dimension(plain, tmp_path / "w.xlsx", "A1:B6")
    assert formula_cell_counts(str(wrong)) == FormulaCounts(formulas=4, without_value=4)
    assert _independent_formula_counts(wrong) == (4, 4)


# -- the skills refuse before reading, and write nothing ---------------------------------


@pytest.mark.parametrize("skill_class", [ExcelDeduplicateSkill, ExcelRemoveEmptyRowsSkill])
@pytest.mark.parametrize(
    ("make", "message"),
    [
        (fixtures.merged_cell_sheet, _MERGED),
        (fixtures.header_on_row_three_sheet, _NO_HEADER),
    ],
)
def test_the_skills_refuse_a_sheet_that_isnt_a_plain_table(
    tmp_path: Path, skill_class: type, make, message: str  # noqa: ANN001
) -> None:
    path = make(tmp_path / "uploads" / "report.xlsx")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    task = SkillTask(
        task_id="t",
        session_id="s",
        workspace_id="w",
        description="clean it",
        input_references=[str(path)],
        output_dir=str(tmp_path / "out"),
    )
    result = skill_class().execute(task)

    assert result.status is SkillStatus.FAILURE
    assert result.errors == [message]
    assert result.outputs == []
    assert not (tmp_path / "out").exists()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


# -- formula counts ----------------------------------------------------------------------


def test_the_formula_counts_match_a_raw_read_of_the_xml(tmp_path: Path) -> None:
    saved = fixtures.formatted_formula_sheet(tmp_path / "saved.xlsx", saved=True)
    unsaved = fixtures.formatted_formula_sheet(tmp_path / "unsaved.xlsx", saved=False)
    partly = fixtures.save_formula_values(unsaved, tmp_path / "partly.xlsx", {"C2": 20, "C6": 60})
    plain = fixtures.plain_sheet(tmp_path / "plain.xlsx")

    assert formula_cell_counts(str(saved)) == FormulaCounts(formulas=4, without_value=0)
    assert formula_cell_counts(str(unsaved)) == FormulaCounts(formulas=4, without_value=4)
    assert formula_cell_counts(str(partly)) == FormulaCounts(formulas=4, without_value=2)
    assert formula_cell_counts(str(plain)) == FormulaCounts(formulas=0, without_value=0)
    for path in (saved, unsaved, partly, plain):
        counts = formula_cell_counts(str(path))
        assert (counts.formulas, counts.without_value) == _independent_formula_counts(path)


def test_the_formula_counts_cover_every_sheet(tmp_path: Path) -> None:
    path = fixtures.formatted_formula_sheet(tmp_path / "r.xlsx", saved=False)
    book = openpyxl.load_workbook(path)
    book.create_sheet("Other")["A1"] = "=1+1"
    book.save(path)
    assert formula_cell_counts(str(path)) == FormulaCounts(formulas=5, without_value=5)
    assert _independent_formula_counts(path) == (5, 5)


def test_an_empty_string_result_counts_as_no_saved_value(tmp_path: Path) -> None:
    unsaved = fixtures.formatted_formula_sheet(tmp_path / "u.xlsx", saved=False)
    path = fixtures.save_formula_values(unsaved, tmp_path / "r.xlsx", {"C2": 20})
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    crafted = xml.replace('<c r="C3"><f>B3*2</f><v />', '<c r="C3" t="str"><f>B3*2</f><v></v>')
    assert crafted != xml
    edits = {"xl/worksheets/sheet1.xml": crafted.encode()}
    out = fixtures.rewrite(path, tmp_path / "e.xlsx", edits)
    assert formula_cell_counts(str(out)) == FormulaCounts(formulas=4, without_value=3)
