"""The PDF table extraction through the whole pipeline, and its independent verifier.

Everything runs the real ``build_engine`` (real Brain, Think, Planner, Hands, the sandboxed PDF
worker, Inspector and Voice) with only the LLM faked. The delivered file is read back with
openpyxl, and the PDF's hash is compared before and after every run (rule 2).
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from collections.abc import Callable
from pathlib import Path

import openpyxl
import pytest

from buddy import engine_factory
from buddy.actions import (
    ACTION_VOCABULARY,
    ACTIONS_BY_CAPABILITY,
    ACTIONS_BY_TARGET,
    EXTRACT_PDF_TABLE,
    INTENT_CAPABILITIES,
    mentioned_actions,
)
from buddy.brain import INTENT_VOCABULARY, ClaudeInterpreter
from buddy.engine_factory import build_engine
from buddy.inspector import PdfTableVerifier, VerificationRequest
from buddy.models import (
    EXTRACTED_TABLE_TARGET,
    ExecutionResult,
    StepResult,
    Understanding,
)
from buddy.outcome import TABLE_TEXT_SENTENCE, VALUES_ONLY_SENTENCE
from buddy.state import BuddyState
from buddy.voice import _COMPLETION_CLAIM, PersonalityVoice
from skills.base import SkillResult, SkillStatus, SkillTask
from skills.pdf import PdfExtractTableSkill, sandbox
from skills.registry import SkillRegistry
from tests import pdf_fixtures as fixtures
from tests.buddy.test_engine_factory import (
    _chain_replies,
    _QueuedLLMProvider,
    _single_intent_replies,
)

pytestmark = pytest.mark.skipif(
    not sys.platform.startswith("linux"),
    reason="the sandbox is Linux-only (it fails closed elsewhere)",
)

SPEAKERS = ("template", "friend", "focus")


@pytest.fixture(autouse=True)
def _work_in_a_temporary_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Skills write under ``workspace/outputs/`` relative to the working folder."""
    monkeypatch.chdir(tmp_path)


EXTRACT_MESSAGE = "extract the table from fruit.pdf"


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _cells(path: str | Path) -> list[list[object]]:
    book = openpyxl.load_workbook(path)
    try:
        return [[c.value for c in row] for row in book.worksheets[0].iter_rows()]
    finally:
        book.close()


def _run(
    tmp_path: Path,
    pdf: bytes | Callable[[Path], object],
    *,
    message: str = EXTRACT_MESSAGE,
    replies: list[str] | None = None,
    speaker: str = "template",
    name: str = "fruit.pdf",
) -> tuple[BuddyState, Path, str]:
    """One turn on ``uploads/<name>`` through ``build_engine``; the upload's hash beforehand."""
    (tmp_path / "uploads").mkdir(exist_ok=True)
    path = tmp_path / "uploads" / name
    if callable(pdf):
        pdf(path)
    else:
        path.write_bytes(pdf)
    before = _sha(path)
    state = BuddyState.new(session_id="session-1", workspace_id="workspace-1")
    state.user_message = message
    state.context.add_uploaded_file(str(path))
    state.context.switch_active_file(str(path))
    voice = None
    if speaker != "template":
        state.context.set_mode(speaker)
        voice = PersonalityVoice()
    provider = _QueuedLLMProvider(replies or _single_intent_replies(EXTRACT_PDF_TABLE))
    build_engine(llm_provider=provider, speaker=voice).run(state)
    assert _sha(path) == before, "the PDF changed"
    return state, path, before


# -- The happy path, in every speaker -----------------------------------------------------------


@pytest.mark.parametrize("speaker", SPEAKERS)
def test_a_pdf_table_is_delivered_verified_and_described_by_what_was_measured(
    tmp_path: Path, speaker: str
) -> None:
    state, pdf, _ = _run(tmp_path, lambda p: fixtures.write_table_pdf(p), speaker=speaker)

    result = state.execution_result
    assert result is not None and result.success and len(result.outputs) == 1
    (output,) = result.outputs
    assert output.endswith(".xlsx") and Path(output) != pdf
    assert _cells(output) == fixtures.FRUIT
    verification = state.verification_result
    assert verification is not None and verification.passed
    (verdict,) = [v for v in verification.steps if v.capability == "pdf.extract_table"]
    (measurement,) = verdict.measurements
    assert (measurement.target, measurement.output_rows, measurement.columns) == (
        EXTRACTED_TABLE_TARGET,
        5,
        3,
    )

    reply = state.reply
    assert reply is not None
    assert "Done — extracted the table: 3 columns, a first row and 4 rows below it." in reply
    assert VALUES_ONLY_SENTENCE in reply and TABLE_TEXT_SENTENCE in reply
    assert "cleaned copy" not in reply and "duplicate" not in reply
    assert "pdf.extract_table" not in reply  # no raw capability id
    assert reply.count("Done") == 1  # the lead word, once
    for text in ("Apple", "Pear", "Plum"):
        assert text not in reply  # nothing from the file reaches the reply
    if speaker == "template":
        assert f"Saved as `{Path(output).name}`." in reply
    else:
        assert "The table from fruit.pdf is in a new spreadsheet, ready for you." in reply


def test_the_pdf_file_name_and_the_output_name_do_not_carry_the_session_id_into_the_download() -> (
    None
):
    from backend.names import download_name

    assert download_name("fruit_table_abc123-task-0.xlsx", "abc123") == "fruit_table.xlsx"
    assert download_name("a_very_long_name_" * 5 + "_table_abc123-task-0.xlsx", "abc123").endswith(
        "_table.xlsx"
    )


# -- Brain's closed vocabulary: no word list decides what a message means -------------------------


def test_the_new_action_is_in_the_closed_vocabularies_and_the_one_action_table() -> None:
    assert EXTRACT_PDF_TABLE == "extract_pdf_table"
    assert EXTRACT_PDF_TABLE in INTENT_VOCABULARY and EXTRACT_PDF_TABLE in ACTION_VOCABULARY
    assert INTENT_CAPABILITIES[EXTRACT_PDF_TABLE] == "pdf.extract_table"
    action = ACTIONS_BY_CAPABILITY["pdf.extract_table"]
    assert action.mention is None  # no pattern, so no word list can add it to a ledger
    assert action.extracts_table and not action.removes_rows
    assert EXTRACTED_TABLE_TARGET not in ACTIONS_BY_TARGET  # it isn't a count of rows left


@pytest.mark.parametrize(
    "message",
    [
        "extract the table from the pdf",
        "please pull the table out of fruit.pdf into excel",
        "convert this pdf to a spreadsheet",
        "extract_pdf_table",
    ],
)
def test_no_code_reads_a_message_for_the_pdf_action(message: str) -> None:
    assert mentioned_actions(message) == ()


def test_a_brain_that_does_not_read_an_extraction_gets_none(tmp_path: Path) -> None:
    """Whatever the words, only Brain's reading runs the skill: 'unclear' runs nothing."""
    state, _, _ = _run(
        tmp_path,
        lambda p: fixtures.write_table_pdf(p),
        replies=_single_intent_replies("unclear"),
    )
    assert state.execution_result is not None
    assert state.execution_result.outputs == () and state.execution_result.steps == ()
    assert state.reply is not None and "no changes were made" in state.reply.lower()
    assert _COMPLETION_CLAIM.search(state.reply) is None, state.reply


def test_brains_prompt_lists_the_new_action_and_intent() -> None:
    from buddy.brain import _SYSTEM_PROMPT

    assert _SYSTEM_PROMPT.count('"extract_pdf_table"') >= 3
    assert "table in a PDF" in _SYSTEM_PROMPT
    parsed = ClaudeInterpreter._parse_reply(
        json.dumps(
            {
                "intent": "extract_pdf_table",
                "requested_actions": [
                    {"action": "extract_pdf_table", "source_phrase": "pull the table out"}
                ],
            }
        ),
        "pull the table out of fruit.pdf",
    )
    assert parsed.intent == "extract_pdf_table"
    assert [(e.action, e.source_phrase) for e in parsed.requested_actions] == [
        ("extract_pdf_table", "pull the table out")
    ]


# -- Refusals: a factual reply and no file ----------------------------------------------------


@pytest.mark.parametrize("speaker", SPEAKERS)
@pytest.mark.parametrize(
    ("make", "words"),
    [
        (fixtures.write_scanned_pdf, "no text layer"),
        (fixtures.write_text_only_pdf, "didn't find a table"),
        (fixtures.write_two_tables_pdf, "more than one table"),
        (
            lambda p: fixtures.write_table_pdf(p, spans=[((0, 1), (1, 1))]),
            "merged or spanning cells",
        ),
        (lambda p: fixtures.write_table_pdf(p, password=("x", "y")), "encrypted"),
        (lambda p: fixtures.write_table_pdf(p, text_pages_after=20), "more than 20 pages"),
        (lambda p: p.write_bytes(fixtures.hostile_deep_page_tree()), "couldn't be read"),
    ],
    ids=["scanned", "no-table", "two-tables", "merged", "encrypted", "21-pages", "hostile"],
)
def test_a_refusal_is_factual_hands_back_no_file_and_keeps_the_pdf(
    tmp_path: Path, speaker: str, make: Callable[[Path], object], words: str
) -> None:
    state, _, _ = _run(tmp_path, make, speaker=speaker)

    result, reply = state.execution_result, state.reply
    assert result is not None and reply is not None
    assert result.outputs == () and not result.success
    assert state.verification_result is not None and not state.verification_result.passed
    assert words in reply
    assert "no file was handed back" in reply and "your original file is unchanged" in reply
    assert _COMPLETION_CLAIM.search(reply) is None, reply
    for forbidden in ("cleaned copy", "the copy", "ready", "pdf.extract_table"):
        assert forbidden not in reply.lower()
    outputs = list(tmp_path.rglob("*.xlsx"))
    assert outputs == [], outputs


# -- Requests that don't fit the file ------------------------------------------------------------


def test_removing_duplicates_from_a_pdf_runs_nothing_and_says_so(tmp_path: Path) -> None:
    state, _, _ = _run(
        tmp_path,
        lambda p: fixtures.write_table_pdf(p),
        message="remove the duplicate rows",
        replies=_single_intent_replies("remove_duplicate_rows"),
    )
    assert state.execution_result is not None and state.execution_result.outputs == ()
    assert state.reply is not None and _COMPLETION_CLAIM.search(state.reply) is None
    assert "no file was handed back" in state.reply


def test_extracting_a_pdf_table_from_a_workbook_runs_nothing_and_says_so(tmp_path: Path) -> None:
    import pandas as pd

    def make(path: Path) -> None:
        pd.DataFrame({"a": [1, 1], "b": [2, 2]}).to_excel(path, index=False)

    state, _, _ = _run(tmp_path, make, name="report.xlsx", message="extract the table")
    assert state.execution_result is not None and state.execution_result.outputs == ()
    assert state.reply is not None and _COMPLETION_CLAIM.search(state.reply) is None
    assert "no file was handed back" in state.reply


def test_extract_and_then_clean_in_one_message_does_not_start(tmp_path: Path) -> None:
    """A chain across file types can't run (the next step can't read a PDF): nothing runs."""
    replies = _chain_replies(
        ("pdf.extract_table", "extract the table"),
        ("spreadsheet.deduplicate", "remove the duplicate rows"),
    )
    state, _, _ = _run(
        tmp_path,
        lambda p: fixtures.write_table_pdf(p),
        message="extract the table and then remove the duplicate rows",
        replies=replies,
    )
    result = state.execution_result
    assert result is not None and result.outputs == () and not result.success
    assert [s.status for s in result.steps].count("success") == 0
    assert list(tmp_path.rglob("*.xlsx")) == []  # not even an intermediate file
    reply = state.reply
    assert reply is not None and _COMPLETION_CLAIM.search(reply) is None, reply
    assert "no file was handed back" in reply


# -- The independent verifier ---------------------------------------------------------------------


def _workbook(path: Path, rows: list[list[object]], *, text: bool = True) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Table"
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value in (None, ""):
                continue
            cell = sheet.cell(row=r, column=c)
            cell.value = value
            if text and isinstance(value, str):
                cell.data_type = "s"
    book.save(path)
    return path


def _verify(pdf: Path, workbook: Path | None) -> object:
    outputs = (str(workbook),) if workbook else ()
    step = StepResult(
        capability="pdf.extract_table",
        skill="pdf.extract_table",
        status="success",
        outputs=outputs,
        input_reference=str(pdf),
        permission_tier="write",
    )
    request = VerificationRequest(
        understanding=Understanding(),
        execution_result=ExecutionResult(success=True, outputs=outputs, steps=(step,)),
    )
    return PdfTableVerifier()(request)


@pytest.fixture
def fruit_pdf(tmp_path: Path) -> Path:
    return fixtures.write_table_pdf(tmp_path / "fruit.pdf")


def test_the_verifier_passes_exactly_the_pdfs_table_and_measures_it(
    tmp_path: Path, fruit_pdf: Path
) -> None:
    result = _verify(fruit_pdf, _workbook(tmp_path / "ok.xlsx", fixtures.FRUIT))
    assert result.passed and result.confidence == 1.0  # type: ignore[attr-defined]
    (m,) = result.measurements  # type: ignore[attr-defined]
    assert (m.target, m.output_rows, m.columns, m.removed) == (EXTRACTED_TABLE_TARGET, 5, 3, None)


def test_an_empty_cell_at_the_end_of_a_row_or_a_column_is_still_the_tables_size(
    tmp_path: Path,
) -> None:
    rows = [["A", "B", ""], ["x", "y", ""], ["", "", ""]]
    pdf = tmp_path / "t.pdf"
    pdf.write_bytes(fixtures.hand_table_pdf(rows))
    result = _verify(pdf, _workbook(tmp_path / "w.xlsx", [["A", "B"], ["x", "y"]]))
    assert result.passed  # type: ignore[attr-defined]
    (m,) = result.measurements  # type: ignore[attr-defined]
    assert (m.output_rows, m.columns) == (3, 3)


def _one_wrong(rows: list[list[str]], r: int, c: int, value: object) -> list[list[object]]:
    out: list[list[object]] = [list(row) for row in rows]
    out[r][c] = value
    return out


@pytest.mark.parametrize(
    ("name", "mutate", "finding"),
    [
        (
            "one wrong cell",
            lambda: _one_wrong(fixtures.FRUIT, 2, 1, "4"),
            "1 cell of the spreadsheet differs",
        ),
        (
            "two wrong cells",
            lambda: _one_wrong(_one_wrong(fixtures.FRUIT, 1, 0, "Pear"), 4, 2, "9"),
            "2 cells",
        ),
        (
            "a swapped row",
            lambda: [fixtures.FRUIT[0], fixtures.FRUIT[2], fixtures.FRUIT[1], *fixtures.FRUIT[3:]],
            "differ",
        ),
        (
            "a number for a text cell",
            lambda: _one_wrong(fixtures.FRUIT, 1, 1, 10),
            "isn't plain text",
        ),
        ("a formula", lambda: _one_wrong(fixtures.FRUIT, 1, 1, "=1+1"), "isn't plain text"),
        (
            "an extra row",
            lambda: [*fixtures.FRUIT, ["Fig", "1", "1.00"]],
            "larger than the PDF's table",
        ),
        (
            "an extra column",
            lambda: [[*row, "x"] for row in fixtures.FRUIT],
            "larger than the PDF's table",
        ),
        ("a dropped row", lambda: fixtures.FRUIT[:-1], "differ"),
        ("an extra space", lambda: _one_wrong(fixtures.FRUIT, 1, 0, "Apple "), "differs"),
    ],
)
def test_the_verifier_fails_unless_every_cell_matches(
    tmp_path: Path, fruit_pdf: Path, name: str, mutate: Callable[[], list], finding: str
) -> None:
    rows = mutate()
    text = name != "a formula"  # a "=" string written the normal way becomes a real formula
    result = _verify(fruit_pdf, _workbook(tmp_path / "w.xlsx", rows, text=text))
    assert not result.passed, name  # type: ignore[attr-defined]
    assert result.measurements == ()  # type: ignore[attr-defined]
    assert finding in " ".join(result.findings), result.findings  # type: ignore[attr-defined]
    # A finding counts what is wrong; it never quotes a cell.
    assert not any(w in " ".join(result.findings) for w in ("Apple", "Plum", "Pear"))  # type: ignore[attr-defined]


def test_a_workbook_with_a_second_sheet_or_another_name_fails(
    tmp_path: Path, fruit_pdf: Path
) -> None:
    ok = _workbook(tmp_path / "ok.xlsx", fixtures.FRUIT)
    book = openpyxl.load_workbook(ok)
    book.create_sheet("Extra")
    book.save(tmp_path / "two.xlsx")
    assert not _verify(fruit_pdf, tmp_path / "two.xlsx").passed  # type: ignore[attr-defined]
    book = openpyxl.load_workbook(ok)
    book.active.title = "Sheet1"
    book.save(tmp_path / "named.xlsx")
    assert not _verify(fruit_pdf, tmp_path / "named.xlsx").passed  # type: ignore[attr-defined]


@pytest.mark.parametrize(
    ("make", "finding"),
    [
        (fixtures.write_two_tables_pdf, "found 2 tables, not exactly one"),
        (fixtures.write_text_only_pdf, "found 0 tables"),
        (lambda p: fixtures.write_table_pdf(p, spans=[((0, 1), (1, 1))]), "isn't a plain grid"),
        (lambda p: fixtures.write_table_pdf(p, password=("x", "y")), "encrypted"),
    ],
)
def test_the_verifier_fails_when_the_second_reading_finds_another_kind_of_pdf(
    tmp_path: Path, make: Callable[[Path], object], finding: str
) -> None:
    pdf = tmp_path / "other.pdf"
    make(pdf)
    result = _verify(pdf, _workbook(tmp_path / "w.xlsx", fixtures.FRUIT))
    assert not result.passed and finding in " ".join(result.findings)  # type: ignore[attr-defined]


def test_the_verifier_fails_when_it_cannot_read_what_it_must(
    tmp_path: Path, fruit_pdf: Path
) -> None:
    ok = _workbook(tmp_path / "ok.xlsx", fixtures.FRUIT)
    assert not _verify(tmp_path / "missing.pdf", ok).passed  # type: ignore[attr-defined]
    assert not _verify(fruit_pdf, tmp_path / "missing.xlsx").passed  # type: ignore[attr-defined]
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a workbook")
    assert not _verify(fruit_pdf, junk).passed  # type: ignore[attr-defined]
    assert not _verify(fruit_pdf, None).passed  # type: ignore[attr-defined]


def test_the_verifier_does_not_trust_what_the_skill_said_about_itself(
    tmp_path: Path, fruit_pdf: Path
) -> None:
    """The step's summary and findings never reach it: a step that claims all fails anyway."""
    wrong = _workbook(tmp_path / "w.xlsx", _one_wrong(fixtures.FRUIT, 1, 1, "99"))
    step = StepResult(
        capability="pdf.extract_table",
        skill="pdf.extract_table",
        status="success",
        outputs=(str(wrong),),
        summary="Extracted a table of 5 rows and 3 columns; every cell checked.",
        self_reported_findings=("all cells match",),
        input_reference=str(fruit_pdf),
        permission_tier="write",
    )
    request = VerificationRequest(
        Understanding(), ExecutionResult(success=True, outputs=(str(wrong),), steps=(step,))
    )
    assert not PdfTableVerifier()(request).passed


def test_the_final_file_check_has_no_pdf_to_compare_and_checks_the_sheets_own_shape(
    tmp_path: Path,
) -> None:
    ok = _workbook(tmp_path / "ok.xlsx", fixtures.FRUIT)
    step = StepResult("pdf.extract_table", "pdf.extract_table", "success", (str(ok),))
    request = VerificationRequest(
        Understanding(), ExecutionResult(success=True, outputs=(str(ok),), steps=(step,))
    )
    assert PdfTableVerifier()(request).passed
    formula = _workbook(tmp_path / "f.xlsx", [["a"], ["=1"]], text=False)
    bad = StepResult("pdf.extract_table", "pdf.extract_table", "success", (str(formula),))
    assert not PdfTableVerifier()(
        VerificationRequest(
            Understanding(), ExecutionResult(success=True, outputs=(str(formula),), steps=(bad,))
        )
    ).passed


# A parser that the skill and its check both use would make them agree on a wrong answer. The
# worker's modes are run here and the libraries each imported are listed.
_WHICH_PARSERS = """
import json, sys
sys.path.append({worker_dir!r})
import main
main.{mode}(open({pdf!r}, 'rb').read())
print(json.dumps(sorted(m for m in sys.modules if m.split('.')[0] in
      ('pdfplumber', 'pdfminer', 'pypdfium2', 'pypdf'))))
"""


@pytest.mark.parametrize(
    ("mode", "uses", "never"),
    [("extract", "pdfminer", "pypdfium2"), ("verify", "pypdfium2", "pdfminer")],
)
def test_the_skill_and_its_check_use_different_pdf_libraries(
    tmp_path: Path, mode: str, uses: str, never: str
) -> None:
    import subprocess

    pdf = fixtures.write_table_pdf(tmp_path / "a.pdf")
    script = _WHICH_PARSERS.format(worker_dir=str(sandbox.WORKER.parent), mode=mode, pdf=str(pdf))
    done = subprocess.run(
        [sys.executable, "-I", "-c", script], capture_output=True, text=True, check=True, timeout=60
    )
    loaded = {m.split(".")[0] for m in json.loads(done.stdout.strip().splitlines()[-1])}
    assert uses in loaded
    assert never not in loaded, f"{mode} must not load {never}: {sorted(loaded)}"
    assert "pdfplumber" not in loaded or mode == "extract"


# An engine whose skill says it succeeded and wrote a table that isn't the PDF's.
class _LyingExtractSkill(PdfExtractTableSkill):
    def __init__(self, tamper: Callable[[Path], None]) -> None:
        self._tamper = tamper

    def execute(self, task: SkillTask) -> SkillResult:
        result = super().execute(task)
        if result.status is SkillStatus.SUCCESS:
            self._tamper(Path(result.outputs[0]))
        return result


def _with_skill(monkeypatch: pytest.MonkeyPatch, skill: PdfExtractTableSkill) -> None:
    original = engine_factory._build_registry

    def build() -> SkillRegistry:
        registry = original()
        registry.remove("pdf.extract_table")
        registry.register(skill)
        return registry

    monkeypatch.setattr(engine_factory, "_build_registry", build)


def _change_a_cell(output: Path) -> None:
    book = openpyxl.load_workbook(output)
    book.worksheets[0]["B3"].value = "4"
    book.save(output)


def _make_a_number(output: Path) -> None:
    book = openpyxl.load_workbook(output)
    book.worksheets[0]["B2"].value = 10
    book.save(output)


def _drop_a_row(output: Path) -> None:
    book = openpyxl.load_workbook(output)
    book.worksheets[0].delete_rows(3)
    book.save(output)


@pytest.mark.parametrize("speaker", SPEAKERS)
@pytest.mark.parametrize(
    "tamper",
    [_change_a_cell, _make_a_number, _drop_a_row],
    ids=["wrong-cell", "number", "dropped-row"],
)
def test_a_skill_that_writes_a_table_that_isnt_the_pdfs_is_stopped_by_the_inspector(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, tamper: Callable[[Path], None], speaker: str
) -> None:
    _with_skill(monkeypatch, _LyingExtractSkill(tamper))
    state, _, _ = _run(tmp_path, lambda p: fixtures.write_table_pdf(p), speaker=speaker)

    result = state.execution_result
    assert result is not None and result.outputs == () and not result.success
    assert state.verification_result is not None and not state.verification_result.passed
    reply = state.reply
    assert reply is not None
    assert "no file was handed back" in reply and "your original file is unchanged" in reply
    assert _COMPLETION_CLAIM.search(reply) is None, reply
    assert "Extracted the table" not in reply and "ready" not in reply


def test_a_skill_that_writes_nothing_is_stopped_by_its_own_wellformedness_check(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Empty(PdfExtractTableSkill):
        def execute(self, task: SkillTask) -> SkillResult:
            path = Path(task.output_dir or ".") / "empty.xlsx"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"")
            return SkillResult(SkillStatus.SUCCESS, outputs=[str(path)])

    _with_skill(monkeypatch, Empty())
    state, _, _ = _run(tmp_path, lambda p: fixtures.write_table_pdf(p))
    assert state.execution_result is not None and state.execution_result.outputs == ()
    assert state.reply is not None and "no file was handed back" in state.reply


def test_the_pdf_is_read_only_for_the_skill_even_when_the_upload_is_writable(
    tmp_path: Path,
) -> None:
    """The run leaves the PDF's bytes and its modification time exactly as they were."""
    (tmp_path / "uploads").mkdir()
    path = fixtures.write_table_pdf(tmp_path / "uploads" / "fruit.pdf")
    stamp = os.stat(path).st_mtime_ns
    state = BuddyState.new("s", "w")
    state.user_message = EXTRACT_MESSAGE
    state.context.add_uploaded_file(str(path))
    state.context.switch_active_file(str(path))
    build_engine(llm_provider=_QueuedLLMProvider(_single_intent_replies("extract_pdf_table"))).run(
        state
    )
    assert os.stat(path).st_mtime_ns == stamp
    assert re.search(r"Done — extracted the table", state.reply or "")


# -- Which file a PDF request reads ------------------------------------------------------------


def test_a_request_to_extract_a_pdfs_table_reads_the_pdf_not_the_copy_handed_back() -> None:
    from buddy.models import FILE_KIND_COPY, FILE_KIND_ORIGINAL, FileChoice
    from buddy.source import choose_file

    choice = FileChoice(
        original="uploads/fruit.pdf", copy="work/fruit_table.xlsx", copy_name="t.xlsx"
    )

    def used(intent: str):  # type: ignore[no-untyped-def]
        state = BuddyState.new("s", "w")
        state.file_choice = choice
        state.context.add_uploaded_file("uploads/fruit.pdf")
        choose_file(state, Understanding(intent=intent))
        return state.file_used

    extraction = used(EXTRACT_PDF_TABLE)
    assert extraction is not None
    assert (extraction.kind, extraction.path) == (FILE_KIND_ORIGINAL, "uploads/fruit.pdf")
    assert extraction.copy_skipped and not extraction.start_over
    cleaning = used("remove_duplicate_rows")
    assert cleaning is not None and cleaning.kind == FILE_KIND_COPY  # the default is unchanged
