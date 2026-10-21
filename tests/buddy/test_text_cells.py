"""A text cell stays text: "007" and 7 are different cells in the cleaning skills, their verifiers
and the analysis.

pandas' default read turns a column of text that looks like numbers into numbers. The cleaning
skills, the Inspector's independent checks and the analysis' duplicate checks read with
``dtype=object`` so they all count rows the same way, and a skill that did convert would be caught.
(Salvaged from ``feat/pdf-extract``, without any PDF code.)
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from buddy import engine_factory
from buddy.engine_factory import build_engine
from buddy.inspector import SpreadsheetAnalysisVerifier, VerificationRequest
from buddy.models import ExecutionResult, StepResult, Understanding
from buddy.state import BuddyState
from skills.base import SkillResult, SkillTask
from skills.excel import ExcelDeduplicateSkill, ExcelRemoveEmptyRowsSkill, SpreadsheetAnalysisSkill
from tests.buddy.test_engine_factory import _QueuedLLMProvider, _single_intent_replies


def _workbook(path: Path, rows: list[list[object]]) -> Path:
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    for row in book.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # stays text even if it looks like a number
    book.save(path)
    return path


DIFFERENT_CELLS = [
    [["id", "n"], ["007", "a"], ["7", "a"]],  # text 007 and text 7 are different cells
    [["id", "n"], ["7", "a"], [7, "a"]],  # text 7 and the number 7 are different cells
    [["x", "y"], ["2.50", "a"], ["2.5", "a"]],
]
DIFFERENT_IDS = ["leading-zero", "text-and-number", "prices"]


@pytest.mark.parametrize("rows", DIFFERENT_CELLS, ids=DIFFERENT_IDS)
def test_the_analysis_does_not_call_two_different_cells_an_exact_duplicate(
    tmp_path: Path, rows: list[list[object]]
) -> None:
    path = _workbook(tmp_path / "w.xlsx", rows)
    task = SkillTask("t", "s", "w", "d", input_references=[str(path)])
    result = SpreadsheetAnalysisSkill().execute(task)
    assert result.metadata["finding_count"] == 0, result.summary
    # The Inspector's independent count agrees with the skill's.
    step = StepResult(
        "spreadsheet.analyze", "x", "success", input_reference=str(path), permission_tier="read"
    )
    verdict = SpreadsheetAnalysisVerifier()(
        VerificationRequest(Understanding(), ExecutionResult(True, steps=(step,)), str(path))
    )
    assert not any("earlier row exactly" in f for f in verdict.findings)
    assert not any("once spaces" in f for f in verdict.findings)


@pytest.mark.parametrize("rows", DIFFERENT_CELLS, ids=DIFFERENT_IDS)
def test_dedupe_keeps_rows_whose_cells_differ_in_type_and_keeps_text_as_text(
    tmp_path: Path, rows: list[list[object]]
) -> None:
    path = _workbook(tmp_path / "w.xlsx", rows)
    out = ExcelDeduplicateSkill().execute(
        SkillTask("t2", "s", "w", "d", input_references=[str(path)], output_dir=str(tmp_path / "o"))
    )
    assert len(pd.read_excel(out.outputs[0], dtype=object)) == len(rows) - 1
    first = openpyxl.load_workbook(out.outputs[0]).active["A2"].value
    assert first == rows[1][0] and type(first) is type(rows[1][0])


def test_a_true_duplicate_of_text_cells_is_still_removed(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "w.xlsx", [["id", "n"], ["007", "a"], ["007", "a"], ["8", "b"]])
    out = ExcelDeduplicateSkill().execute(
        SkillTask("t3", "s", "w", "d", input_references=[str(path)], output_dir=str(tmp_path / "o"))
    )
    rows = list(openpyxl.load_workbook(out.outputs[0]).active.iter_rows(values_only=True))
    assert rows == [("id", "n"), ("007", "a"), ("8", "b")]


def test_empty_row_removal_keeps_text_cells_as_text(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "w.xlsx", [["id", "n"], ["007", "a"], [None, None], ["008", "b"]])
    out = ExcelRemoveEmptyRowsSkill().execute(
        SkillTask("t4", "s", "w", "d", input_references=[str(path)], output_dir=str(tmp_path / "o"))
    )
    rows = list(openpyxl.load_workbook(out.outputs[0]).active.iter_rows(values_only=True))
    assert rows == [("id", "n"), ("007", "a"), ("008", "b")]


def _converting(skill_class: type) -> type:
    """A cleaning skill that, after running, rewrites the text ``007`` in its output as 7."""

    class Converting(skill_class):  # type: ignore[misc, valid-type]
        def execute(self, task: SkillTask) -> SkillResult:
            result = super().execute(task)
            for output in result.outputs:
                book = openpyxl.load_workbook(output)
                for row in book.active.iter_rows():
                    for cell in row:
                        if cell.value == "007":
                            cell.value = 7
                book.save(output)
            return result

    return Converting


@pytest.mark.parametrize(
    ("skill_class", "skill_name", "intent"),
    [
        (ExcelDeduplicateSkill, "excel.deduplicate", "remove_duplicate_rows"),
        (ExcelRemoveEmptyRowsSkill, "excel.remove_empty_rows", "remove_empty_rows"),
    ],
)
def test_a_cleaning_step_that_turns_text_into_numbers_fails_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type, skill_name: str, intent: str
) -> None:
    original = engine_factory._build_registry

    def build():  # type: ignore[no-untyped-def]
        registry = original()
        registry.remove(skill_name)
        registry.register(_converting(skill_class)())
        return registry

    monkeypatch.setattr(engine_factory, "_build_registry", build)
    (tmp_path / "uploads").mkdir()
    path = _workbook(
        tmp_path / "uploads" / "ids.xlsx",
        [["id", "n"], ["007", "a"], ["008", "b"], ["008", "b"], [None, None], ["009", "c"]],
    )
    state = BuddyState.new("session-1", "workspace-1")
    state.user_message = "clean it"
    state.context.add_uploaded_file(str(path))
    state.context.switch_active_file(str(path))
    build_engine(llm_provider=_QueuedLLMProvider(_single_intent_replies(intent))).run(state)
    assert state.execution_result is not None and state.execution_result.outputs == ()
    assert state.verification_result is not None and not state.verification_result.passed
    assert "no file was handed back" in (state.reply or "")
