"""Backlog item 1a (review M3): multi-sheet workbooks, end to end through ``build_engine``.

The row-cleaning skills read and write one sheet, so they refuse a workbook
with more than one: nothing is delivered and the reply says why. The
verifier also compares each file-changing step's sheet names with its
input's, so a skill that drops or renames a sheet fails verification
whatever it reported. All LLM calls are faked.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from buddy.engine_factory import build_engine
from buddy.outcome import VALUES_ONLY_SENTENCE
from buddy.state import BuddyState
from buddy.voice import _COMPLETION_CLAIM, PersonalityVoice
from buddy.voice_catalog import WORKED_ON_ORIGINAL_SENTENCE
from skills.base import SkillResult, SkillStatus
from skills.excel.dedupe import ExcelDeduplicateSkill
from skills.excel.remove_empty_rows import ExcelRemoveEmptyRowsSkill
from tests import workbook_fixtures as fixtures
from tests.buddy.test_engine_factory import (
    _chain_replies,
    _QueuedLLMProvider,
    _single_intent_replies,
)

_Q1 = {"name": ["Alice", "Bob", None, "Alice", "Carol"], "amount": [1, 2, None, 1, 3]}
_Q2 = {"name": ["Zed", "Yan"], "amount": [9, 8]}

_UPLOAD = Path("uploads") / "report.xlsx"

_DEDUPE = "spreadsheet.deduplicate"
_EMPTY_ROWS = "spreadsheet.remove_empty_rows"

# kind -> (message, the replies Brain/Think(/ClaudePlanner) get, the file-changing steps)
_REQUESTS = {
    "dedupe": (
        "remove the duplicate rows",
        lambda: _single_intent_replies("remove_duplicate_rows"),
        [_DEDUPE],
    ),
    "empty_rows": (
        "remove the empty rows",
        lambda: _single_intent_replies("remove_empty_rows"),
        [_EMPTY_ROWS],
    ),
    "chain": (
        "remove the duplicate rows and then remove the empty rows",
        lambda: _chain_replies(
            (_DEDUPE, "remove the duplicate rows"), (_EMPTY_ROWS, "remove the empty rows")
        ),
        [_DEDUPE, _EMPTY_ROWS],
    ),
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sheet_names(path: str | Path) -> list[str]:
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def _write_workbook(path: Path, sheets: dict[str, dict[str, list[object]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path) as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, index=False)


def _sheets(count: int) -> dict[str, dict[str, list[object]]]:
    """``Q1`` (a duplicate and an empty row), then ``count - 1`` more sheets."""
    return {"Q1": _Q1, **{f"Q{i}": _Q2 for i in range(2, count + 1)}}


def _run(
    tmp_path: Path, sheets: dict[str, dict[str, list[object]]], kind: str, speaker_kind: str
) -> tuple[BuddyState, _QueuedLLMProvider, Path, str]:
    """One turn through ``build_engine`` on ``uploads/report.xlsx``."""
    _write_workbook(tmp_path / _UPLOAD, sheets)
    return _run_uploaded(tmp_path, kind, speaker_kind)


def _run_uploaded(
    tmp_path: Path, kind: str, speaker_kind: str
) -> tuple[BuddyState, _QueuedLLMProvider, Path, str]:
    """One turn through ``build_engine`` on the file already at ``uploads/report.xlsx``."""
    message, replies, _ = _REQUESTS[kind]
    input_path = _UPLOAD
    original_hash = _sha256(tmp_path / input_path)

    state = BuddyState.new(session_id="session-1", workspace_id="workspace-1")
    state.user_message = message
    state.context.add_uploaded_file(str(input_path))
    state.context.switch_active_file(str(input_path))
    speaker = None
    if speaker_kind != "template":
        state.context.set_mode(speaker_kind)
        speaker = PersonalityVoice()
    provider = _QueuedLLMProvider(replies())
    build_engine(llm_provider=provider, speaker=speaker).run(state)
    return state, provider, tmp_path / input_path, original_hash


def _assert_nothing_delivered(state: BuddyState) -> str:
    """Nothing handed back, nothing claimed; returns the reply."""
    result = state.execution_result
    assert result is not None
    assert result.outputs == ()
    assert state.verification_result is not None
    assert state.verification_result.passed is False
    reply = state.reply
    assert reply is not None
    assert "no file was handed back" in reply
    assert "your original file is unchanged" in reply
    lowered = reply.lower()
    assert "cleaned copy" not in lowered
    assert "ready" not in lowered
    assert "the copy" not in lowered
    assert _COMPLETION_CLAIM.search(reply) is None, reply
    assert "spreadsheet." not in reply  # no raw capability id
    return reply


# -- The skills refuse a multi-sheet workbook -------------------------------------------


@pytest.mark.parametrize("sheet_count", [2, 3])
@pytest.mark.parametrize("kind", list(_REQUESTS))
@pytest.mark.parametrize("speaker_kind", ["template", "friend", "focus"])
def test_a_multi_sheet_workbook_is_refused_and_nothing_is_delivered(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    sheet_count: int,
    kind: str,
    speaker_kind: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    state, provider, original_path, original_hash = _run(
        tmp_path, _sheets(sheet_count), kind, speaker_kind
    )
    steps = _REQUESTS[kind][2]

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == steps[0]
    first = next(s for s in result.steps if s.capability == steps[0])
    assert first.status == "failure"
    # Every later step never ran: step 2 would have read step 1's output.
    assert [s.status for s in result.steps if s.capability in steps[1:]] == ["skipped"] * (
        len(steps) - 1
    )

    reply = _assert_nothing_delivered(state)
    assert (
        f"the workbook has {sheet_count} sheets, and only single-sheet workbooks are "
        "supported for now" in reply
    )
    # The refusal is the only reason given: no "no output file" leftover.
    assert "did not produce an output file" not in reply

    # Nothing was written, the original is byte-identical, and the Voice model
    # was never called (every queued reply was used by Brain/Think/planner).
    assert not list((tmp_path / "workspace").rglob("*.xlsx"))
    assert _sha256(original_path) == original_hash
    assert _sheet_names(original_path) == [f"Q{i}" for i in range(1, sheet_count + 1)]
    assert len(provider.calls) == len(_REQUESTS[kind][1]())


# -- A one-sheet workbook behaves exactly as before -------------------------------------

# The replies ``main`` (310f660) gives for this fixture, captured with the same
# faked LLM replies. Only the sheet name of the output changed: it's now kept. Since
# ``fix/values-only-copy`` every delivered copy also says what it doesn't keep, right after
# the file sentence (``_with_values_only_sentence``).
_MAIN_REPLIES = {
    ("dedupe", "template"): (
        "Done — removed 1 duplicate row. The copy has 4 rows, with 0 duplicate rows "
        "remaining. Saved as `report_deduplicated_session-1-task-1.xlsx`. I also noticed: "
        "name: 1 missing value (25.0%); amount: 1 missing value (25.0%). Want me to "
        "tackle any of these?"
    ),
    ("dedupe", "friend"): (
        "Okay, here's the rundown. Done — removed 1 duplicate row. The copy has 4 rows, "
        "with 0 duplicate rows remaining. Your cleaned copy of report.xlsx is ready. I also "
        "noticed: name: 1 missing value (25.0%); amount: 1 missing value (25.0%). Want "
        "me to tackle any of these?"
    ),
    ("dedupe", "focus"): (
        "Done — removed 1 duplicate row. The copy has 4 rows, with 0 duplicate rows "
        "remaining. Your cleaned copy of report.xlsx is ready. I also noticed: name: 1 "
        "missing value (25.0%); amount: 1 missing value (25.0%). Want me to tackle any "
        "of these?"
    ),
    ("chain", "template"): (
        "Removed 1 duplicate row. Removed 1 fully-empty row. The copy has 3 rows, with 0 "
        "duplicate rows and 0 fully-empty rows remaining. Saved as "
        "`report_deduplicated_session-1-task-1_no_empty_rows_session-1-task-2.xlsx`. None of "
        "the issues I check for were found."
    ),
    ("chain", "friend"): (
        "Those blank rows were just taking up space. Removed 1 duplicate row. Removed 1 "
        "fully-empty row. The copy has 3 rows, with 0 duplicate rows and 0 fully-empty "
        "rows remaining. Your cleaned copy of report.xlsx is ready. None of the issues I "
        "check for were found."
    ),
    ("chain", "focus"): (
        "Removed 1 duplicate row. Removed 1 fully-empty row. The copy has 3 rows, with 0 "
        "duplicate rows and 0 fully-empty rows remaining. Your cleaned copy of "
        "report.xlsx is ready. None of the issues I check for were found."
    ),
}


def _with_values_only_sentence(reply: str) -> str:
    """``reply`` with the values-only sentence after its file sentence, then the sentence saying
    which file the turn worked on (the person's upload: no copy was offered)."""
    file_sentence = re.search(
        r"Saved as `[^`]+`\.|Your cleaned copy of report\.xlsx is ready\.", reply
    )
    assert file_sentence is not None
    end = file_sentence.end()
    return f"{reply[:end]} {VALUES_ONLY_SENTENCE} {WORKED_ON_ORIGINAL_SENTENCE}{reply[end:]}"


@pytest.mark.parametrize("sheet_name", ["Sheet1", "Data"])
@pytest.mark.parametrize(("kind", "speaker_kind"), list(_MAIN_REPLIES))
def test_a_one_sheet_workbook_behaves_exactly_as_before(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    sheet_name: str,
    kind: str,
    speaker_kind: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    state, _, original_path, original_hash = _run(
        tmp_path, {sheet_name: _Q1}, kind, speaker_kind
    )

    result = state.execution_result
    assert result is not None
    assert result.stopped_at is None
    assert len(result.outputs) == 1
    assert state.verification_result is not None
    assert state.verification_result.passed is True
    assert state.reply == _with_values_only_sentence(_MAIN_REPLIES[(kind, speaker_kind)])
    # The one sheet keeps its name (``main`` renamed it to "Sheet1").
    assert _sheet_names(result.outputs[0]) == [sheet_name]
    assert _sha256(original_path) == original_hash


# -- The verifier's structural net: a skill that drops or renames a sheet ---------------


def _first_sheet_only(input_path: str, output_path: Path) -> None:
    """What the skills did on ``main``: read the first sheet, write ``Sheet1``."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pd.read_excel(input_path).drop_duplicates().to_excel(output_path, index=False)


def _drops_every_other_sheet(self, task):  # type: ignore[no-untyped-def]
    """A fake dedupe with ``main``'s behaviour: no refusal, one sheet out."""
    output_path = self._build_output_path(task)
    _first_sheet_only(task.input_references[0], output_path)
    return SkillResult(status=SkillStatus.SUCCESS, outputs=[str(output_path)])


@pytest.mark.parametrize("speaker_kind", ["template", "friend", "focus"])
def test_a_skill_that_drops_a_sheet_fails_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, speaker_kind: str
) -> None:
    monkeypatch.setattr(ExcelDeduplicateSkill, "execute", _drops_every_other_sheet)
    monkeypatch.chdir(tmp_path)
    state, _, original_path, original_hash = _run(tmp_path, _sheets(2), "dedupe", speaker_kind)

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == _DEDUPE
    assert any(e.startswith(f"{_DEDUPE} failed verification") for e in result.errors)
    dedupe = next(s for s in state.verification_result.steps if s.capability == _DEDUPE)
    assert dedupe.passed is False
    assert dedupe.findings == (
        "the output's sheets (Sheet1) don't match its input's (Q1, Q2), so a sheet was "
        "lost or renamed",
    )

    reply = _assert_nothing_delivered(state)
    assert "so a sheet was lost or renamed" in reply
    assert _sha256(original_path) == original_hash


def test_a_skill_that_renames_the_only_sheet_fails_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(ExcelDeduplicateSkill, "execute", _drops_every_other_sheet)
    monkeypatch.chdir(tmp_path)
    state, _, original_path, original_hash = _run(tmp_path, {"Data": _Q1}, "dedupe", "template")

    assert state.execution_result is not None
    assert state.execution_result.stopped_at == _DEDUPE
    dedupe = next(s for s in state.verification_result.steps if s.capability == _DEDUPE)
    assert dedupe.findings == (
        "the output's sheets (Sheet1) don't match its input's (Data), so a sheet was lost "
        "or renamed",
    )
    _assert_nothing_delivered(state)
    assert _sha256(original_path) == original_hash


def _keeps_every_sheet(self, task):  # type: ignore[no-untyped-def]
    """A fake dedupe that dedupes the first sheet and copies the others."""
    input_path = task.input_references[0]
    output_path = self._build_output_path(task)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheets = pd.read_excel(input_path, sheet_name=None)
    first = next(iter(sheets))
    with pd.ExcelWriter(output_path) as writer:
        for name, frame in sheets.items():
            frame = frame.drop_duplicates() if name == first else frame
            frame.to_excel(writer, sheet_name=name, index=False)
    return SkillResult(status=SkillStatus.SUCCESS, outputs=[str(output_path)])


def _drops_every_other_sheet_empty_rows(self, task):  # type: ignore[no-untyped-def]
    """A fake empty-row removal with ``main``'s behaviour: one sheet out."""
    output_path = self._build_output_path(task)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pd.read_excel(task.input_references[0]).dropna(how="all").to_excel(output_path, index=False)
    return SkillResult(status=SkillStatus.SUCCESS, outputs=[str(output_path)])


@pytest.mark.parametrize("speaker_kind", ["template", "friend", "focus"])
def test_a_chain_step_that_drops_a_sheet_of_the_previous_steps_output_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, speaker_kind: str
) -> None:
    """Step 1 keeps both sheets and passes. Step 2 reads step 1's output and
    drops ``Q2``: its sheet names are checked against *that* file, so it
    fails, the chain stops, and nothing is handed back."""
    monkeypatch.setattr(ExcelDeduplicateSkill, "execute", _keeps_every_sheet)
    monkeypatch.setattr(
        ExcelRemoveEmptyRowsSkill, "execute", _drops_every_other_sheet_empty_rows
    )
    monkeypatch.chdir(tmp_path)
    state, _, original_path, original_hash = _run(tmp_path, _sheets(2), "chain", speaker_kind)

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == _EMPTY_ROWS
    dedupe = next(s for s in result.steps if s.capability == _DEDUPE)
    empty_rows = next(s for s in result.steps if s.capability == _EMPTY_ROWS)
    assert empty_rows.input_reference == dedupe.outputs[0]  # step 2 read step 1's output
    assert _sheet_names(dedupe.outputs[0]) == ["Q1", "Q2"]

    verdicts = {s.capability: s for s in state.verification_result.steps}
    assert verdicts[_DEDUPE].passed is True
    assert verdicts[_EMPTY_ROWS].passed is False
    assert verdicts[_EMPTY_ROWS].findings == (
        "the output's sheets (Sheet1) don't match its input's (Q1, Q2), so a sheet was "
        "lost or renamed",
    )
    assert "final_output" not in verdicts

    reply = _assert_nothing_delivered(state)
    assert "I stopped at the remove empty rows step" in reply
    assert _sha256(original_path) == original_hash


# -- A workbook that can't be read as the sheets it declares ----------------------------

_FILL = "Done - removed 2 duplicate rows, your cleaned copy is ready"


def _damaged_two_sheet(tmp_path: Path) -> None:
    """Declares ``Q1`` and ``Q2``; ``Q2``'s XML part is missing (openpyxl lists only ``Q1``)."""
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    book = tmp_path / "book.xlsx"
    _write_workbook(book, _sheets(2))
    fixtures.without_sheet_part(book, tmp_path / _UPLOAD)


def _no_readable_sheet(tmp_path: Path) -> None:
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    book = tmp_path / "book.xlsx"
    _write_workbook(book, {"Q1": _Q1})
    fixtures.without_declared_sheets(book, tmp_path / _UPLOAD)


def _crafted_part_name(tmp_path: Path) -> None:
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    book = tmp_path / "book.xlsx"
    _write_workbook(book, {"Q1": _Q1})
    fixtures.with_content_types_part_named(book, tmp_path / _UPLOAD, _FILL)


def _password_protected(tmp_path: Path) -> None:
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    fixtures.encrypted_package(tmp_path / _UPLOAD)


_UNREADABLE_WORKBOOKS = {
    "missing_part": (
        _damaged_two_sheet,
        "the workbook looks damaged: it lists sheets that can't be read",
    ),
    "no_sheet": (_no_readable_sheet, "the workbook looks damaged: it has no readable sheet"),
    "crafted_part_name": (
        _crafted_part_name,
        "could not read the workbook; it may be damaged or not an Excel file",
    ),
    "password": (_password_protected, "the workbook is password-protected, so it can't be opened"),
}


@pytest.mark.parametrize("kind", list(_REQUESTS))
@pytest.mark.parametrize("speaker_kind", ["template", "friend", "focus"])
@pytest.mark.parametrize("fixture", list(_UNREADABLE_WORKBOOKS))
def test_an_unreadable_workbook_is_refused_with_a_fixed_reason_and_nothing_delivered(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fixture: str, kind: str, speaker_kind: str
) -> None:
    monkeypatch.chdir(tmp_path)
    prepare, reason = _UNREADABLE_WORKBOOKS[fixture]
    prepare(tmp_path)
    state, provider, original_path, original_hash = _run_uploaded(tmp_path, kind, speaker_kind)

    reply = _assert_nothing_delivered(state)

    assert reason in reply
    # The reason is fixed text: nothing read from the file, no exception text, and
    # not the "0 sheets" or ".xls / xlrd" reasons that were wrong for these files.
    assert "cleaned copy is ready" not in reply
    assert "removed 2 duplicate rows" not in reply
    assert "0 sheets" not in reply
    assert "xlrd" not in reply
    assert ".xls" not in reply
    assert "did not produce an output file" not in reply
    assert not list((tmp_path / "workspace").rglob("*.xlsx"))
    assert _sha256(original_path) == original_hash
    assert len(provider.calls) == len(_REQUESTS[kind][1]())  # the Voice model never ran


def test_a_crafted_part_name_with_a_completion_claim_never_reaches_the_reply(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    _crafted_part_name(tmp_path)
    state, _, _, _ = _run_uploaded(tmp_path, "dedupe", "template")

    reply = state.reply
    assert reply is not None
    assert "cleaned copy is ready" not in reply
    assert "removed 2 duplicate rows" not in reply
    assert _COMPLETION_CLAIM.search(reply) is None, reply


# -- The proactive analysis is skipped on a workbook the skills refuse (M4) -------------


def _capabilities(state: BuddyState) -> list[str]:
    assert state.execution_result is not None
    return [step.capability for step in state.execution_result.steps]


@pytest.mark.parametrize("kind", list(_REQUESTS))
@pytest.mark.parametrize("speaker_kind", ["template", "friend", "focus"])
@pytest.mark.parametrize("sheet_count", [2, 3])
def test_the_analysis_is_skipped_on_a_multi_sheet_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    sheet_count: int,
    kind: str,
    speaker_kind: str,
) -> None:
    monkeypatch.chdir(tmp_path)
    state, _, _, _ = _run(tmp_path, _sheets(sheet_count), kind, speaker_kind)

    assert "spreadsheet.analyze" not in _capabilities(state)
    reply = _assert_nothing_delivered(state)
    # The reply describes no single sheet and offers no work the skills refuse.
    assert "here's what I noticed" not in reply
    assert "I also noticed" not in reply
    assert "tackle" not in reply
    assert "missing value" not in reply
    assert f"the workbook has {sheet_count} sheets" in reply
    # Skipped, so the file wasn't marked as analysed either.
    assert str(_UPLOAD) not in state.context.analyzed_files


def test_a_hidden_first_sheet_is_not_described(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Sheet 1 is hidden: the old analysis described a sheet the user can't see."""
    monkeypatch.chdir(tmp_path)
    _write_workbook(tmp_path / _UPLOAD, _sheets(2))
    book = openpyxl.load_workbook(tmp_path / _UPLOAD)
    book["Q1"].sheet_state = "hidden"
    book.active = 1
    book.save(tmp_path / _UPLOAD)
    state, _, _, _ = _run_uploaded(tmp_path, "dedupe", "template")

    assert "spreadsheet.analyze" not in _capabilities(state)
    reply = _assert_nothing_delivered(state)
    assert "missing value" not in reply


@pytest.mark.parametrize("fixture", list(_UNREADABLE_WORKBOOKS))
def test_the_analysis_is_skipped_on_a_workbook_whose_sheets_cant_be_listed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fixture: str
) -> None:
    monkeypatch.chdir(tmp_path)
    _UNREADABLE_WORKBOOKS[fixture][0](tmp_path)
    state, _, _, _ = _run_uploaded(tmp_path, "dedupe", "template")

    assert "spreadsheet.analyze" not in _capabilities(state)


@pytest.mark.parametrize("kind", list(_REQUESTS))
def test_the_analysis_still_runs_on_a_one_sheet_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    monkeypatch.chdir(tmp_path)
    state, _, _, _ = _run(tmp_path, {"Data": _Q1}, kind, "template")

    assert _capabilities(state)[0] == "spreadsheet.analyze"
    assert str(_UPLOAD) in state.context.analyzed_files
