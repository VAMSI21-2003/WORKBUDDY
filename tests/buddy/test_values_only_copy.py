"""Finding P-1: a delivered copy holds values only, and a sheet that isn't a plain table is refused.

End to end through ``build_engine`` (only the LLM faked), all three speakers, a single step and
a chain, on: a formatted sheet with formulas (saved and unsaved), a sheet with merged cells, a
header that isn't on the first row, and a plain sheet. Every number in a reply is checked
against a read of the files that uses neither pandas nor the skills' or the Inspector's code,
and the original's hash is checked before and after. ``docs/design/values-only-copy.md``.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pytest

from buddy import inspector
from buddy.models import OriginalFacts
from buddy.outcome import (
    FORMULAS_UNCHECKED_SENTENCE,
    VALUES_ONLY_SENTENCE,
    _values_only_sentences,
)
from skills.excel.dedupe import ExcelDeduplicateSkill
from skills.excel.remove_empty_rows import ExcelRemoveEmptyRowsSkill
from skills.excel.workbook import single_sheet_name
from tests import workbook_fixtures as fixtures
from tests.buddy.test_multi_sheet import (
    _REQUESTS,
    _assert_nothing_delivered,
    _run_uploaded,
    _sha256,
)

_UPLOAD = Path("uploads") / "report.xlsx"
_SUFFIX = ", and only plain tables (one header row at the top, then the data) are supported for now"
_MERGED = "the sheet has merged cells" + _SUFFIX
_NO_HEADER = "a column in the sheet has no header" + _SUFFIX
_READS_AS_EMPTY = (
    "a cell holds an error value, or text such as NA or N/A, that would be read as empty, "
    "and keeping such cells isn't supported for now"
)

_SPEAKERS = ["template", "friend", "focus"]
_KINDS = list(_REQUESTS)  # dedupe, empty_rows, chain

# name -> (builder, formula cells, of those with no saved value). ``None``: not a plain table.
_DELIVERED = {
    "formulas_saved": (lambda p: fixtures.formatted_formula_sheet(p, saved=True), 4, 0),
    "formulas_unsaved": (lambda p: fixtures.formatted_formula_sheet(p, saved=False), 4, 4),
    "plain": (fixtures.plain_sheet, 0, 0),
}
_REFUSED = {
    "merged_cells": (fixtures.merged_cell_sheet, _MERGED),
    "header_on_row_3": (fixtures.header_on_row_three_sheet, _NO_HEADER),
    "na_text": (fixtures.na_text_sheet, _READS_AS_EMPTY),
    "formula_error_results": (fixtures.error_result_sheet, _READS_AS_EMPTY),
    "header_gap_wrong_dimension": (fixtures.gap_header_wrong_dimension_sheet, _NO_HEADER),
}


# -- an independent read of the files ----------------------------------------------------


def _rows(path: Path) -> list[tuple[object, ...]]:
    """Every row below the header of the first sheet, as saved values (openpyxl, not pandas)."""
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        return [tuple(row) for row in book.worksheets[0].iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()


def _blank(row: tuple[object, ...]) -> bool:
    return all(value is None or value == "" for value in row)


def _formula_cells(path: Path) -> tuple[int, int]:
    """Formula cells, and those with no saved value, from the sheet XML itself."""
    formulas = unsaved = 0
    with zipfile.ZipFile(path) as archive:
        sheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    for cell in sheet.iter():
        if cell.tag.endswith("}c") and any(child.tag.endswith("}f") for child in cell):
            formulas += 1
            value = next((c for c in cell if c.tag.endswith("}v")), None)
            unsaved += value is None or not (value.text or "")
    return formulas, unsaved


def _expected(original: Path, kind: str) -> tuple[int, int, int]:
    """(duplicate rows removed, empty rows removed, rows in the copy), worked out by hand."""
    rows = _rows(original)
    unique = list(dict.fromkeys(rows))
    if kind == "dedupe":
        return len(rows) - len(unique), 0, len(unique)
    if kind == "empty_rows":
        kept = [row for row in rows if not _blank(row)]
        return 0, len(rows) - len(kept), len(kept)
    kept = [row for row in unique if not _blank(row)]
    return len(rows) - len(unique), len(unique) - len(kept), len(kept)


def _number(pattern: str, reply: str) -> int | None:
    found = re.findall(pattern, reply)
    assert len(found) <= 1, (pattern, reply)
    return int(found[0]) if found else None


def _formulas_sentence(count: int) -> str:
    if count == 1:
        return "Your file has 1 formula cell. The copy has its last saved value, not the formula."
    return (
        f"Your file has {count} formula cells. The copy has their last saved values, "
        "not the formulas."
    )


def _unsaved_sentence(count: int, formulas: int = 3) -> str:
    if formulas == 1:
        return (
            "That cell had no saved value, so it's empty in the copy, and rows were judged "
            "with it empty."
        )
    if count == 1:
        return (
            "1 of them had no saved value, so it's empty in the copy, and rows were judged "
            "with that cell empty."
        )
    return (
        f"{count} of them had no saved value, so they're empty in the copy, and rows were "
        "judged with those cells empty."
    )


# -- a plain table is delivered, and the reply says what the copy doesn't keep -------------


@pytest.mark.parametrize("speaker_kind", _SPEAKERS)
@pytest.mark.parametrize("kind", _KINDS)
@pytest.mark.parametrize("fixture", list(_DELIVERED))
def test_a_delivered_copy_says_what_it_doesnt_keep(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fixture: str, kind: str, speaker_kind: str
) -> None:
    monkeypatch.chdir(tmp_path)
    build, formulas, unsaved = _DELIVERED[fixture]
    original = build(tmp_path / _UPLOAD)
    # The fixture's own numbers are an independent read, not an assumption.
    assert _formula_cells(original) == (formulas, unsaved)

    state, provider, original_path, original_hash = _run_uploaded(tmp_path, kind, speaker_kind)

    result = state.execution_result
    assert result is not None
    assert result.stopped_at is None
    assert len(result.outputs) == 1
    assert state.verification_result is not None
    assert state.verification_result.passed is True
    reply = state.reply
    assert reply is not None
    delivered = Path(result.outputs[0])

    # The numbers in the reply agree with the files.
    removed_dupes, removed_empty, copy_rows = _expected(original_path, kind)
    assert _number(r"[Rr]emoved (\d+) duplicate row", reply) == (removed_dupes or None)
    assert _number(r"[Rr]emoved (\d+) fully-empty row", reply) == (removed_empty or None)
    assert _number(r"The copy has (\d+) rows?", reply) == copy_rows
    assert len(_rows(delivered)) == copy_rows
    assert _number(r"Your file has (\d+) formula cell", reply) == (formulas or None)
    if formulas != 1:
        assert _number(r"(\d+) of them had no saved value", reply) == (unsaved or None)

    # The sentences are there, once, in order, right after the file sentence.
    file_sentence = (
        "Saved as `" if speaker_kind == "template" else "Your cleaned copy of report.xlsx is ready."
    )
    assert reply.count(VALUES_ONLY_SENTENCE) == 1
    assert reply.index(file_sentence) < reply.index(VALUES_ONLY_SENTENCE)
    if formulas:
        assert reply.count(_formulas_sentence(formulas)) == 1
        assert reply.index(VALUES_ONLY_SENTENCE) < reply.index(_formulas_sentence(formulas))
    else:
        assert "formula" not in reply
    if unsaved:
        assert reply.count(_unsaved_sentence(unsaved, formulas)) == 1
    else:
        assert "no saved value" not in reply

    # ...and what they say is true of the delivered file.
    assert _formula_cells(delivered)[0] == 0  # "not the formulas"
    copy = openpyxl.load_workbook(delivered)
    try:
        sheet = copy.worksheets[0]
        assert sheet.column_dimensions["A"].width != 40  # a width isn't carried over
        assert sheet["B2"].fill.fgColor.rgb != "FFFFFF00"  # nor a fill
    finally:
        copy.close()
    if fixture == "formulas_unsaved":
        values = _rows(delivered)
        assert all(row[2] is None for row in values)  # "empty in the copy"
    if fixture == "formulas_saved":
        assert {row[2] for row in _rows(delivered) if row[0] is not None} <= {20, 40, 60}

    # The Voice model wrote nothing, and the original is byte-identical.
    assert len(provider.calls) == len(_REQUESTS[kind][1]())
    assert _sha256(original_path) == original_hash


@pytest.mark.parametrize("speaker_kind", _SPEAKERS)
def test_a_plain_sheet_reply_is_the_old_reply_plus_the_values_only_sentence(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, speaker_kind: str
) -> None:
    monkeypatch.chdir(tmp_path)
    fixtures.plain_sheet(tmp_path / _UPLOAD)
    state, _, _, _ = _run_uploaded(tmp_path, "dedupe", speaker_kind)

    reply = state.reply
    assert reply is not None
    assert reply.startswith(
        {
            "template": (
                "Done — removed 1 duplicate row. The copy has 4 rows, with 0 duplicate rows "
                "remaining. Saved as `report_deduplicated_session-1-task-1.xlsx`. "
            ),
            "friend": (
                "Okay, here's the rundown. Done — removed 1 duplicate row. The copy has 4 "
                "rows, with 0 duplicate rows remaining. Your cleaned copy of report.xlsx is "
                "ready. "
            ),
            "focus": (
                "Done — removed 1 duplicate row. The copy has 4 rows, with 0 duplicate rows "
                "remaining. Your cleaned copy of report.xlsx is ready. "
            ),
        }[speaker_kind]
        + VALUES_ONLY_SENTENCE
    )


# -- a sheet that isn't a plain table is refused, and nothing is delivered ------------------


@pytest.mark.parametrize("speaker_kind", _SPEAKERS)
@pytest.mark.parametrize("kind", _KINDS)
@pytest.mark.parametrize("fixture", list(_REFUSED))
def test_a_sheet_that_isnt_a_plain_table_is_refused_and_nothing_is_delivered(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fixture: str, kind: str, speaker_kind: str
) -> None:
    monkeypatch.chdir(tmp_path)
    build, message = _REFUSED[fixture]
    build(tmp_path / _UPLOAD)
    state, provider, original_path, original_hash = _run_uploaded(tmp_path, kind, speaker_kind)
    steps = _REQUESTS[kind][2]

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == steps[0]
    first = next(s for s in result.steps if s.capability == steps[0])
    assert first.status == "failure"
    assert [s.status for s in result.steps if s.capability in steps[1:]] == ["skipped"] * (
        len(steps) - 1
    )

    reply = _assert_nothing_delivered(state)
    assert message in reply
    # None of the old false claims: no "Done", no copy, no Unnamed column from the misread
    # sheet, no analysis of a sheet the skills refuse, no word about formulas or formatting.
    assert "Unnamed" not in reply
    assert "I also noticed" not in reply
    assert VALUES_ONLY_SENTENCE not in reply
    assert "formula" not in reply

    assert not list((tmp_path / "workspace").rglob("*.xlsx"))
    assert _sha256(original_path) == original_hash
    assert len(provider.calls) == len(_REQUESTS[kind][1]())


# -- the Inspector checks it independently of the skills ------------------------------------


def _skills_dont_refuse(monkeypatch: pytest.MonkeyPatch) -> None:
    """Skills with ``main``'s behaviour: only the sheet count is checked, so they deliver."""
    monkeypatch.setattr("skills.excel.dedupe.plain_table_sheet_name", single_sheet_name)
    monkeypatch.setattr("skills.excel.remove_empty_rows.plain_table_sheet_name", single_sheet_name)
    assert ExcelDeduplicateSkill and ExcelRemoveEmptyRowsSkill  # both modules are imported


@pytest.mark.parametrize("speaker_kind", _SPEAKERS)
@pytest.mark.parametrize("kind", _KINDS)
@pytest.mark.parametrize("fixture", list(_REFUSED))
def test_a_skill_that_doesnt_refuse_still_fails_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fixture: str, kind: str, speaker_kind: str
) -> None:
    _skills_dont_refuse(monkeypatch)
    monkeypatch.chdir(tmp_path)
    build, message = _REFUSED[fixture]
    build(tmp_path / _UPLOAD)
    state, _, original_path, original_hash = _run_uploaded(tmp_path, kind, speaker_kind)
    steps = _REQUESTS[kind][2]

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == steps[0]
    assert state.verification_result is not None
    first = next(s for s in state.verification_result.steps if s.capability == steps[0])
    assert first.passed is False
    assert first.findings == (f"the sheet it read isn't a plain table: {message}",)

    reply = _assert_nothing_delivered(state)
    assert "the sheet it read isn't a plain table" in reply
    assert message in reply
    assert "Unnamed" not in reply
    assert _sha256(original_path) == original_hash


def test_the_check_on_a_chain_step_reads_the_file_that_step_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Step 1 is fine; step 2 reads step 1's output, which a skill that writes a merged cell
    # would make non-plain. The check on step 2's input catches it.
    monkeypatch.chdir(tmp_path)
    fixtures.plain_sheet(tmp_path / _UPLOAD)
    real_write = ExcelDeduplicateSkill.execute

    def writes_a_merged_cell(self, task):  # type: ignore[no-untyped-def]
        outcome = real_write(self, task)
        book = openpyxl.load_workbook(outcome.outputs[0])
        book.active.merge_cells("E1:F1")
        book.save(outcome.outputs[0])
        return outcome

    monkeypatch.setattr(ExcelDeduplicateSkill, "execute", writes_a_merged_cell)
    state, _, original_path, original_hash = _run_uploaded(tmp_path, "chain", "template")

    result = state.execution_result
    assert result is not None
    assert result.stopped_at == "spreadsheet.remove_empty_rows"
    reply = _assert_nothing_delivered(state)
    assert _MERGED in reply
    assert _sha256(original_path) == original_hash


# -- the measurement, and what happens when it can't be made ---------------------------------


@pytest.mark.parametrize("speaker_kind", _SPEAKERS)
def test_a_failed_read_of_the_original_says_it_couldnt_check(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, speaker_kind: str
) -> None:
    def cant_read(path: str):  # type: ignore[no-untyped-def]
        raise RuntimeError("secret-path-and-text")

    monkeypatch.setattr(inspector, "formula_cell_counts", cant_read)
    monkeypatch.chdir(tmp_path)
    fixtures.formatted_formula_sheet(tmp_path / _UPLOAD, saved=True)
    state, _, original_path, original_hash = _run_uploaded(tmp_path, "dedupe", speaker_kind)

    assert state.verification_result is not None
    assert state.verification_result.passed is True
    assert state.verification_result.original_facts == OriginalFacts(formula_cells=None)
    reply = state.reply
    assert reply is not None
    assert VALUES_ONLY_SENTENCE in reply
    assert FORMULAS_UNCHECKED_SENTENCE in reply
    assert "(s)" not in reply
    assert "secret" not in reply
    assert _sha256(original_path) == original_hash


def test_the_measurement_is_the_original_even_in_a_chain(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Step 2 reads step 1's output, which has no formulas; the count is the user's file's.
    monkeypatch.chdir(tmp_path)
    fixtures.formatted_formula_sheet(tmp_path / _UPLOAD, saved=False)
    state, _, _, _ = _run_uploaded(tmp_path, "chain", "template")
    assert state.verification_result is not None
    assert state.verification_result.original_facts == OriginalFacts(
        formula_cells=4, formula_cells_without_value=4
    )


def test_no_measurement_is_made_for_a_turn_that_delivered_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    fixtures.merged_cell_sheet(tmp_path / _UPLOAD)
    state, _, _, _ = _run_uploaded(tmp_path, "dedupe", "template")
    assert state.verification_result is not None
    assert state.verification_result.original_facts is None


# -- the renderer ----------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("facts", "expected"),
    [
        (None, [VALUES_ONLY_SENTENCE]),
        (OriginalFacts(formula_cells=0), [VALUES_ONLY_SENTENCE]),
        (OriginalFacts(formula_cells=None), [VALUES_ONLY_SENTENCE, FORMULAS_UNCHECKED_SENTENCE]),
        (OriginalFacts(formula_cells=3), [VALUES_ONLY_SENTENCE, _formulas_sentence(3)]),
        (
            OriginalFacts(formula_cells=3, formula_cells_without_value=2),
            [VALUES_ONLY_SENTENCE, _formulas_sentence(3), _unsaved_sentence(2)],
        ),
    ],
)
def test_the_sentences_come_from_the_measurement_and_nothing_else(
    facts: OriginalFacts | None, expected: list[str]
) -> None:
    assert _values_only_sentences(facts) == expected


def test_unsaved_formulas_are_not_mentioned_when_there_are_no_formulas() -> None:
    # A stray count with no formulas states nothing.
    facts = OriginalFacts(formula_cells=0, formula_cells_without_value=5)
    assert _values_only_sentences(facts) == [VALUES_ONLY_SENTENCE]

