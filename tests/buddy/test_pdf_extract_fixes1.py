"""Tests for the findings of review round 1 of feat/pdf-extract.

The findings are in ``docs/reviews/feat-pdf-extract-review.md``.

Each one fails on the code the reviewer read (head e841088). Same setting as ``test_pdf_extract``:
the real pipeline and the real sandboxed worker, only the LLM faked.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from buddy.engine_factory import build_engine
from buddy.inspector import SpreadsheetAnalysisVerifier, VerificationRequest
from buddy.models import ExecutionResult, StepResult, Understanding
from buddy.state import BuddyState
from buddy.voice import _COMPLETION_CLAIM
from skills.base import PermissionTier, SkillResult, SkillTask
from skills.excel import ExcelDeduplicateSkill, ExcelRemoveEmptyRowsSkill, SpreadsheetAnalysisSkill
from skills.pdf import PdfExtractTableSkill, sandbox
from skills.pdf.extract_table import REFUSALS
from tests import pdf_fixtures as fixtures
from tests.buddy.test_engine_factory import (
    _chain_replies,
    _QueuedLLMProvider,
    _single_intent_replies,
)
from tests.buddy.test_pdf_extract import _cells, _run

pytestmark = pytest.mark.skipif(
    not sys.platform.startswith("linux"),
    reason="the sandbox is Linux-only (it fails closed elsewhere)",
)


@pytest.fixture(autouse=True)
def _work_in_a_temporary_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)


def _task(pdf: Path, out: Path) -> SkillTask:
    return SkillTask(
        task_id="t-0",
        session_id="s",
        workspace_id="w",
        description="extract",
        input_references=[str(pdf)],
        permission=PermissionTier.WRITE,
        output_dir=str(out),
    )


# -- H1: text that runs past its cell ---------------------------------------------------------

OVERFLOW_LAST = [["Name", "Comment"], ["Alice", "Paid in full on delivery"], ["Bob", "ok"]]
OVERFLOW_MIDDLE = [["Name", "Note", "Qty"], ["Alice", "a fairly long note that will not fit", ""]]
OVERFLOW_ACCENTS = [["Name", "Text"], ["Zoe", "caf\u00e9 na\u00efve and some more words"]]


@pytest.mark.parametrize(
    "rows", [OVERFLOW_LAST, OVERFLOW_MIDDLE, OVERFLOW_ACCENTS], ids=["last", "middle", "accents"]
)
def test_a_table_with_text_running_past_its_cell_is_refused_not_delivered_cut(
    tmp_path: Path, rows: list[list[str]]
) -> None:
    """Both readings clip a cell to its box the same way, so their agreement proves nothing."""
    pdf = tmp_path / "wide.pdf"
    pdf.write_bytes(fixtures.hand_table_pdf(rows, cell_width=90))
    # The skill, alone, would deliver the cut text (that is the finding)...
    extracted = sandbox.run("extract", pdf.read_bytes())
    assert extracted.ok and extracted.data["rows"] != rows
    # ...so the independent reading refuses the table...
    assert sandbox.run("verify", pdf.read_bytes()).reason == "text_overflows"
    # ...and the turn hands nothing back, and says why in a fixed sentence.
    state, _, _ = _run(tmp_path, lambda p: p.write_bytes(pdf.read_bytes()))
    assert state.execution_result is not None and state.execution_result.outputs == ()
    reply = state.reply or ""
    assert REFUSALS["text_overflows"] in reply
    assert "no file was handed back" in reply
    assert _COMPLETION_CLAIM.search(reply) is None, reply


def test_text_that_fits_its_cells_is_still_delivered(tmp_path: Path) -> None:
    pdf = tmp_path / "ok.pdf"
    pdf.write_bytes(fixtures.hand_table_pdf(fixtures.FRUIT))
    assert sandbox.run("verify", pdf.read_bytes()).ok
    state, _, _ = _run(tmp_path, lambda p: p.write_bytes(pdf.read_bytes()))
    assert state.execution_result is not None and len(state.execution_result.outputs) == 1


def test_text_beside_the_table_on_one_of_its_rows_is_refused(tmp_path: Path) -> None:
    rows = [["a", "b"], ["x", "y"]]
    data = fixtures.hand_table_pdf(rows).replace(
        b"BT /F1 10 Tf 54 ", b"BT /F1 10 Tf 300 ", 1  # the first cell's text, moved off the table
    )
    assert sandbox.run("verify", data).reason == "text_overflows"


# -- H2: the text of a cell is not rewritten into another value ------------------------------


def test_fractions_units_and_superscripts_come_through_unchanged(tmp_path: Path) -> None:
    rows = [["Size", "Area", "Power"], ["2½ in", "12 m²", "10³ W"], ["¼", "m³", "x"]]
    pdf = tmp_path / "units.pdf"
    pdf.write_bytes(fixtures.hand_table_pdf(rows))
    state, _, _ = _run(tmp_path, lambda p: p.write_bytes(pdf.read_bytes()))
    assert state.execution_result is not None and len(state.execution_result.outputs) == 1
    assert _cells(state.execution_result.outputs[0]) == rows


def test_a_trademark_sign_is_not_turned_into_letters(tmp_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    table = Table([["Brand", "Note"], ["Acme™", "ok"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    pdf = tmp_path / "tm.pdf"
    SimpleDocTemplate(str(pdf)).build([table])
    result = PdfExtractTableSkill().execute(_task(pdf, tmp_path / "out"))
    assert result.outputs, result.errors
    assert _cells(result.outputs[0])[1][0] == "Acme™"


# -- H3: no exception's text in a reply ---------------------------------------------------------


def test_a_report_asked_with_the_extraction_never_shows_an_exceptions_words(tmp_path: Path) -> None:
    replies = _chain_replies(
        ("pdf.extract_table", "extract the table"), ("spreadsheet.analyze", "give me a report")
    )
    state, _, _ = _run(
        tmp_path,
        lambda p: fixtures.write_table_pdf(p),
        message="extract the table and give me a report",
        replies=replies,
        speaker="friend",
    )
    reply = state.reply or ""
    assert state.execution_result is not None
    for leak in ("No skill supporting", "-task-", "session-1", "can handle task"):
        assert leak not in reply, reply
    for error in state.execution_result.errors:
        assert "No skill supporting" not in error and "-task-" not in error


# -- M1: the counts in the two replies add up ---------------------------------------------------
# (also in tests/backend/test_pdf_turns.py, which runs the two turns through the API)


def test_the_extraction_sentence_counts_the_first_row_apart_from_the_rows_below_it() -> None:
    from buddy.models import StepMeasurement
    from buddy.outcome import _extraction_sentence

    def sentence(rows: int, columns: int) -> str:
        measurement = StepMeasurement("extracted_table", rows, 0, columns=columns)
        return _extraction_sentence("table extraction", measurement)

    assert sentence(5, 3) == "Extracted the table: 3 columns, a first row and 4 rows below it."
    assert sentence(2, 2) == "Extracted the table: 2 columns, a first row and 1 row below it."
    assert (
        sentence(1, 4) == "Extracted the table: 4 columns and a first row, with no rows below it."
    )


# -- M2: a table ruled around its rows only is not a table of columns -----------------------------


def test_a_table_ruled_only_around_its_rows_is_refused(tmp_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    table = Table([list(row) for row in fixtures.FRUIT], colWidths=[100, 60, 60])
    table.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 1, colors.black),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.black),
            ]
        )
    )
    pdf = tmp_path / "rows-only.pdf"
    SimpleDocTemplate(str(pdf)).build([table])
    result = PdfExtractTableSkill().execute(_task(pdf, tmp_path / "out"))
    assert result.outputs == [] and result.errors == [REFUSALS["one_column"]]
    assert sandbox.run("verify", pdf.read_bytes()).data["plain"] is False
    assert not list((tmp_path / "out").glob("*")) if (tmp_path / "out").exists() else True


# -- M3: the analysis and the cleaning skills count exact duplicates the same way ----------------


def _workbook(path: Path, rows: list[list[object]]) -> Path:
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    for row in book.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    book.save(path)
    return path


@pytest.mark.parametrize(
    "rows",
    [
        [["id", "n"], ["007", "a"], ["7", "a"]],  # text 007 and text 7 are different cells
        [["id", "n"], ["7", "a"], [7, "a"]],  # text 7 and the number 7 are different cells
        [["x", "y"], ["2.50", "a"], ["2.5", "a"]],
    ],
    ids=["leading-zero", "text-and-number", "extracted-prices"],
)
def test_the_analysis_does_not_call_two_different_cells_an_exact_duplicate(
    tmp_path: Path, rows: list[list[object]]
) -> None:
    path = _workbook(tmp_path / "w.xlsx", rows)
    task = SkillTask("t", "s", "w", "d", input_references=[str(path)])
    result = SpreadsheetAnalysisSkill().execute(task)
    # The skill reports a count of what it found; an exact duplicate among them would show here.
    assert result.metadata["finding_count"] == 0, result.summary
    # And the Inspector's independent count agrees with the cleaning skills' (none to remove).
    step = StepResult(
        "spreadsheet.analyze", "x", "success", input_reference=str(path), permission_tier="read"
    )
    verdict = SpreadsheetAnalysisVerifier()(
        VerificationRequest(Understanding(), ExecutionResult(True, steps=(step,)), str(path))
    )
    assert not any("earlier row exactly" in f for f in verdict.findings)
    out = ExcelDeduplicateSkill().execute(
        SkillTask("t2", "s", "w", "d", input_references=[str(path)], output_dir=str(tmp_path / "o"))
    )
    assert len(pd.read_excel(out.outputs[0], dtype=object)) == len(rows) - 1


# -- L8: a cleaning skill that rewrites text cells as numbers is caught -----------------------


class _Converting:
    """Wraps a cleaning skill: runs it, then rewrites the text ``007`` cells of its output as 7."""

    def __init__(self, skill_class: type) -> None:
        self._skill = skill_class()
        for name in ("name", "description", "supported_inputs", "permission_tier", "can_handle"):
            setattr(type(self), name, property(lambda me, n=name: getattr(me._skill, n)))

    def execute(self, task: SkillTask) -> SkillResult:
        result = self._skill.execute(task)
        for output in result.outputs:
            book = openpyxl.load_workbook(output)
            for row in book.active.iter_rows():
                for cell in row:
                    if cell.value == "007":
                        cell.value = 7
            book.save(output)
        return result

    def verify(self, task: SkillTask, result: SkillResult) -> bool:
        return self._skill.verify(task, result)


@pytest.mark.parametrize(
    ("skill_class", "skill_name", "intent"),
    [
        (ExcelDeduplicateSkill, "excel.deduplicate", "remove_duplicate_rows"),
        (ExcelRemoveEmptyRowsSkill, "excel.remove_empty_rows", "remove_empty_rows"),
    ],
)
def test_a_cleaning_step_that_turns_text_into_numbers_fails_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, skill_class: type, skill_name: str, intent: str
) -> None:
    from buddy import engine_factory

    original = engine_factory._build_registry

    def build():  # type: ignore[no-untyped-def]
        registry = original()
        registry.remove(skill_name)
        registry.register(_Converting(skill_class))  # type: ignore[arg-type]
        return registry

    monkeypatch.setattr(engine_factory, "_build_registry", build)
    (tmp_path / "uploads").mkdir()
    path = _workbook(
        tmp_path / "uploads" / "ids.xlsx",
        [["id", "n"], ["007", "a"], ["008", "b"], ["008", "b"], [None, None], ["009", "c"]],
    )
    state = BuddyState.new("session-1", "workspace-1")
    state.user_message = "clean it"
    state.context.add_uploaded_file(str(path))
    state.context.switch_active_file(str(path))
    build_engine(llm_provider=_QueuedLLMProvider(_single_intent_replies(intent))).run(state)
    assert state.execution_result is not None and state.execution_result.outputs == ()
    assert state.verification_result is not None and not state.verification_result.passed
    assert "no file was handed back" in (state.reply or "")


def test_the_worker_cap_is_one_place() -> None:
    """One worker at a time: a dense PDF measured 706 MiB, and the app's container is capped."""
    assert sandbox.MAX_CONCURRENT_WORKERS == 1
    assert sandbox._places._value == 1  # type: ignore[attr-defined]
