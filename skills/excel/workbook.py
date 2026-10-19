"""Workbook — the single-sheet and plain-table checks the row-cleaning skills run first.

``ExcelDeduplicateSkill`` and ``ExcelRemoveEmptyRowsSkill`` read one sheet
and write one sheet, so a workbook with more than one sheet would come back
with every other sheet gone. Until per-sheet cleaning exists, they refuse
such a workbook instead (``docs/backlog.md`` item 1a).

They also write the sheet's *values* back with pandas, so a sheet that isn't
a plain table (one header row, then the data) would come back changed: a
merged note cell as an "Unnamed" column, a header that isn't on the first row
read as data, a duplicate header renamed. They refuse such a sheet too
(``docs/design/values-only-copy.md``). The check reads the file with openpyxl,
never from a skill's DataFrame, and the Inspector runs the same one on each
step's input.

The sheet list comes from ``sheet_names``, which the Inspector uses too: it
reads the sheets *declared* in the workbook's own ``workbook.xml``, then
cross-checks them with openpyxl. A library's list alone isn't trusted,
because openpyxl silently drops a declared sheet whose XML part is missing.
Every message in a ``WorkbookRefused`` is fixed text: it reaches the user's
reply, so it never carries an exception's text or a name read from the file.

A hostile workbook is refused before its table is built: ``check_workbook_limits`` reads the zip
directory (parts, declared unzipped size), refuses any DOCTYPE or entity declaration, and then runs
the workbook through the readers the skills, the verifiers and the analysis use, in a child process
with a memory cap and a time limit (``_scan_in_child``). A 2 MiB ``.xlsx`` can unpack to gigabytes
of XML, a few cells far out in a sheet make a reader pad the grid between them, and openpyxl parses
a dozen kinds of part whole on every load, so no list of parts or elements is complete. What the
child could read inside its limits bounds what the server's own later reads need, because they are
the same reads. A file that needs more, or whose check crashes or times out, is refused.
"""

import json
import logging
import os
import subprocess
import sys
import threading
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree
from xml.parsers import expat

import openpyxl

logger = logging.getLogger(__name__)

# The first bytes of an .xlsx (a zip archive) and of an OLE2 compound file.
# The format is read from the content, not the extension, the same way
# pandas picks its reader.
_ZIP_MAGIC = b"PK\x03\x04"
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Password-protected .xlsx files are OLE2 containers too, holding a stream
# with this name (stored as UTF-16-LE in the container's directory).
_ENCRYPTED_STREAM = "EncryptedPackage".encode("utf-16-le")

# The content types of a workbook's main part: .xlsx and macro-enabled.
_WORKBOOK_CONTENT_TYPES = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
)

# An upper bound on the XML parts read: a real workbook.xml is a few KB.
_MAX_PART_BYTES = 16 * 1024 * 1024

# What a workbook may be before it is refused as too large to open safely. Real tables are far
# below these (the measured 5 MB file, 61,277 rows by 14 columns, is 858,000 cells and about 60 MB
# unpacked); the upload route refuses a declared size above ``MAX_UNZIPPED_BYTES`` too.
MIB = 1024 * 1024
MAX_UNZIPPED_BYTES = 64 * MIB
MAX_PARTS = 1000
MAX_ROWS = 200_000
MAX_COLUMNS = 1_000
MAX_CELLS = 1_000_000
# Excel's own limit for the text in one cell, and a cap on all the text the cells hold.
MAX_CELL_CHARS = 32_767
# The text the cells hold, counted as CPython stores it (1, 2 or 4 bytes a character by the widest
# character in the cell), not in characters: 10 million emoji took 1.9 GiB of a turn.
MAX_TEXT_BYTES = 4_000_000

# The check of what a workbook needs to be opened runs in a child process (``_scan_in_child``)
# under these limits. Unmeasured in the production image until ``scripts/measure_workbook_check.py``
# has been run there: the app container has 1 GiB, and the server plus one child must fit in it.
SCAN_MEMORY_BYTES = 512 * MIB  # address space of the child
SCAN_SECONDS = 60
# How long a check waits for the one running check before saying the server is busy.
SCAN_WAIT_SECONDS = 10

# How much of a sheet's XML is fed to the merged-cell scan at a time.
_SCAN_CHUNK_BYTES = 1 << 20

_UNREADABLE = "could not read the workbook; it may be damaged or not an Excel file"
_NO_SHEET = "the workbook looks damaged: it has no readable sheet"
_SHEETS_UNREADABLE = "the workbook looks damaged: it lists sheets that can't be read"
_PASSWORD = "the workbook is password-protected, so it can't be opened"
_LEGACY_XLS = (
    "older .xls workbooks (Excel 97–2003) aren't supported yet; saving it as .xlsx will work"
)
# Too large to open safely. Fixed text with the limit in it, never a figure read from the file.
TOO_MANY_PARTS = f"the workbook has more than {MAX_PARTS:,} internal parts, which is too many"
TOO_BIG_UNPACKED = (
    f"the workbook unpacks to more than {MAX_UNZIPPED_BYTES // MIB} MB, "
    "which is more than I can open safely"
)
TOO_MANY_ROWS = (
    f"the workbook has more than {MAX_ROWS:,} rows (all its sheets counted together), "
    "which is more than I can open safely"
)
TOO_MANY_COLUMNS = (
    f"the sheet has more than {MAX_COLUMNS:,} columns, which is more than I can open safely"
)
TOO_MANY_CELLS = (
    f"the workbook spans more than {MAX_CELLS:,} cells (each sheet's rows times its columns, "
    "added up), which is more than I can open safely"
)

TOO_LONG_CELL = (
    f"a cell holds more than {MAX_CELL_CHARS:,} characters, which is more than I can open safely"
)
TOO_MUCH_TEXT = (
    f"the cells hold more than {MAX_TEXT_BYTES // 1_000_000} million bytes of text once opened "
    "(emoji and other wide characters count for up to four), which is more than I can open safely"
)
# The child process ran out of memory or time, or ended without an answer. All fixed text: nothing
# the file or the child said reaches it.
TOO_COSTLY = "the workbook needs more memory or time to open than I allow, so I won't open it"
CHECK_STOPPED = (
    "the check of the workbook was stopped before it finished, so I can't tell that it is safe "
    "to open and I won't open it"
)
CHECK_BUSY = (
    "the server is busy checking another workbook, so I couldn't check this one; "
    "please try again in a moment"
)
_NOT_ALLOWED = (
    "the workbook holds XML declarations that an Excel file doesn't use, so I won't open it"
)


# What a sheet that isn't a plain table is refused for. Fixed text: it reaches the
# user's reply, so it never carries a name or a value read from the file.
_MERGED_CELLS = "the sheet has merged cells"
_HEADER_NOT_FIRST = "the sheet's header isn't on the first row"
_COLUMN_WITHOUT_HEADER = "a column in the sheet has no header"
_DUPLICATE_HEADER = "two columns in the sheet have the same header"
_PLAIN_TABLES_ONLY = (
    ", and only plain tables (one header row at the top, then the data) are supported for now"
)
# A cell pandas reads as missing although it isn't empty: an error value (a formula's
# ``#N/A``, a typed ``#DIV/0!``) or text such as ``NA`` or ``null``. A values-only copy
# would hold an empty cell there, and a row of them would look empty to the skills.
_READS_AS_EMPTY = (
    "a cell holds an error value, or text such as NA or N/A, that would be read as empty, "
    "and keeping such cells isn't supported for now"
)
# The strings ``pandas.read_excel`` reads as missing by default (its ``STR_NA_VALUES``,
# less the empty string, which is an empty cell); a test keeps this in step with pandas.
_PANDAS_NA_TEXT = frozenset(
    (
        "-1.#IND", "-1.#QNAN", "-NaN", "-nan", "1.#IND", "1.#QNAN", "<NA>", "N/A", "NA", "NULL",
        "NaN", "None", "n/a", "nan", "null", "#N/A", "#N/A N/A", "#NA",
    )
)


class WorkbookRefused(Exception):
    """The workbook can't be cleaned. ``str(exc)`` is the user-facing reason."""


def single_sheet_name(input_path: str) -> str:
    """Return the name of the workbook's only sheet, opening it read-only.

    Args:
        input_path: The workbook to check. It is only ever read.

    Returns:
        The name of its one sheet.

    Raises:
        WorkbookRefused: It has more than one sheet, or its sheets can't be
            listed. The message says which, in plain words.
    """
    names = sheet_names(input_path)
    if len(names) != 1:
        raise WorkbookRefused(
            f"the workbook has {len(names)} sheets, and only single-sheet workbooks "
            "are supported for now"
        )
    return names[0]


def plain_table_sheet_name(input_path: str) -> str:
    """Return the name of the workbook's only sheet, if it is a plain table.

    The check the row-cleaning skills run before reading anything: one sheet
    (``single_sheet_name``), and that sheet a plain table (``plain_table_problem``).

    Args:
        input_path: The workbook to check. It is only ever read.

    Returns:
        The name of its one sheet.

    Raises:
        WorkbookRefused: It has more than one sheet, can't be read, or its
            sheet isn't a plain table. The message says which, in plain words.
    """
    name = single_sheet_name(input_path)
    problem = _sheet_problem(input_path, name)
    if problem is not None:
        raise WorkbookRefused(problem)
    return name


def plain_table_problem(input_path: str) -> str | None:
    """Say why a sheet of the workbook isn't a plain table, if one isn't.

    A plain table has its header on row 1 with a header over every column up
    to the last, its data below, no merged cells, and no two equal headers:
    the shape ``pandas.read_excel`` and ``to_excel`` round-trip unchanged. A
    sheet with no values at all isn't a problem (there is nothing to change).
    Every sheet the workbook declares is checked, in order.

    Read with openpyxl (read-only), so the verdict is independent of any
    skill's DataFrame; the cell values are the ones a pandas read sees.

    Args:
        input_path: The workbook to check. It is only ever read.

    Returns:
        ``None`` if every sheet is a plain table; otherwise a fixed sentence
        saying which rule the first offending sheet breaks.

    Raises:
        WorkbookRefused: The workbook can't be read (see ``sheet_names``).
    """
    for name in sheet_names(input_path):
        problem = _sheet_problem(input_path, name)
        if problem is not None:
            return problem
    return None


def sheet_names(input_path: str) -> list[str]:
    """Every sheet the workbook declares, worksheets and chart sheets alike, in order.

    Read-only, from the content rather than the extension. For zip content
    (``.xlsx``) the names are the ``<sheet>`` entries of the workbook part,
    and openpyxl must list the same names, or the workbook is refused: a
    declared sheet openpyxl can't load would otherwise be dropped unnoticed.
    OLE2 content (a password-protected ``.xlsx`` or a legacy ``.xls``) is
    refused without being counted.

    Args:
        input_path: The workbook to read. It is only ever read.

    Returns:
        The declared sheet names, at least one.

    Raises:
        WorkbookRefused: The workbook can't be read, is password-protected,
            is a legacy ``.xls``, or looks damaged. The message is fixed text.
    """
    try:
        check_workbook_limits(input_path)
        with open(input_path, "rb") as handle:
            magic = handle.read(8)
            handle.seek(0)
            if magic.startswith(_ZIP_MAGIC):
                return _zip_sheet_names(handle)
            if magic == _OLE2_MAGIC:
                raise WorkbookRefused(_PASSWORD if _has_stream(handle) else _LEGACY_XLS)
    except (WorkbookRefused, MemoryError):
        raise
    except Exception as exc:  # noqa: BLE001 — unreadable or corrupt, reported as a refusal
        logger.exception("could not read the sheet names of %r", input_path)
        raise WorkbookRefused(_UNREADABLE) from exc
    logger.info("%r isn't an Excel workbook", input_path)
    raise WorkbookRefused(_UNREADABLE)


# The verdict for a file already scanned, keyed on what identifies its bytes on disk, so the
# planner, the skill and the Inspector's reads of the same file scan it once. Bounded. A busy
# server's "try again" is not a verdict about the file and is never kept.
_scanned: dict[tuple[str, int, int, int], str | None] = {}
_SCANNED_MAX = 64
_scanned_lock = threading.Lock()
# One child at a time: the server plus one child must fit in the container.
_scan_slot = threading.Semaphore(1)
# Set in the child process, whose own reads must not start another child.
_in_scan_child = False


def check_workbook_limits(input_path: str) -> None:
    """Refuse a workbook that is too large to open safely, before its table is built.

    Only zip content (``.xlsx``) is looked at; anything else is left to the callers' own
    refusals. Checks, each cheaper than the next: the number of parts; the declared unzipped size
    of all parts (``zipfile`` never hands over more of a part than its directory declares, so
    that total bounds what is read); no DOCTYPE or entity declaration in any part; then the
    readers the skills use, run in a child process with a memory cap and a time limit
    (``_scan_in_child``), so that what is counted and what is needed is what the table would be
    built from, whatever the XML says about itself. A file that can't be read is left to the
    callers' own refusals: the table's reader can't read it either. A file whose check ran out of
    memory or time, or crashed, is refused.

    Args:
        input_path: The workbook to check. It is only ever read.

    Raises:
        WorkbookRefused: A limit was passed, or the check could not be run. The message is
            fixed text.
    """
    if _in_scan_child:
        return
    try:
        stat = Path(input_path).stat()
    except OSError:
        return  # the caller's own read says what is wrong with the path
    key = (input_path, stat.st_size, stat.st_mtime_ns, stat.st_ino)
    with _scanned_lock:
        known = key in _scanned
        verdict = _scanned.get(key)
    if not known:
        if not _scan_slot.acquire(timeout=SCAN_WAIT_SECONDS):
            raise WorkbookRefused(CHECK_BUSY)
        try:
            with _scanned_lock:
                known = key in _scanned
                verdict = _scanned.get(key)
            if not known:  # another thread may have checked this very file while this one waited
                verdict = _scan(input_path)
                with _scanned_lock:
                    if len(_scanned) >= _SCANNED_MAX:
                        _scanned.clear()
                    _scanned[key] = verdict
        finally:
            _scan_slot.release()
    if verdict is not None:
        raise WorkbookRefused(verdict)


def _scan(input_path: str) -> str | None:
    """The refusal text for a workbook over a limit, or ``None``."""
    try:
        with open(input_path, "rb") as handle:
            if not handle.read(len(_ZIP_MAGIC)).startswith(_ZIP_MAGIC):
                return None
        with zipfile.ZipFile(input_path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_PARTS:
                return TOO_MANY_PARTS
            if sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES:
                return TOO_BIG_UNPACKED
            # The one part every reader opens first, refused by its declared size as it always
            # was (``_read_part``), before a child is started for it.
            for entry in entries:
                if entry.filename.lower() == "[content_types].xml":
                    if entry.file_size > _MAX_PART_BYTES:
                        return _UNREADABLE
            # Every part, whatever its name: the workbook's own relationships say which ones a
            # reader opens. A part that isn't XML stops the parser at its first bytes.
            for entry in entries:
                if _declares_entities(archive, entry):
                    return _NOT_ALLOWED
    except Exception:  # noqa: BLE001 — damaged or unreadable: the callers' own reads refuse it
        logger.info("%r couldn't be scanned for size limits", input_path)
        return None
    return _scan_in_child(input_path)


def _declares_entities(archive: zipfile.ZipFile, entry: zipfile.ZipInfo) -> bool:
    """Whether the part is XML that declares a DOCTYPE or an entity (read by the parser itself).

    A part that isn't well-formed XML says nothing here: the callers' own reads refuse it.
    """
    parser = expat.ParserCreate()
    found = False

    def declared(*_: object) -> None:
        nonlocal found
        found = True
        raise _LimitPassed

    parser.StartDoctypeDeclHandler = declared
    parser.EntityDeclHandler = declared
    try:
        with archive.open(entry) as part:
            while chunk := part.read(_SCAN_CHUNK_BYTES):
                parser.Parse(chunk, False)
            parser.Parse(b"", True)
    except _LimitPassed:
        pass
    except Exception:  # noqa: BLE001 — not XML, or damaged: not this check's business
        return False
    return found


# What the child can answer, by position: the parent maps the number back to fixed text, so
# nothing the child prints (and nothing from the file) is ever shown.
_CHILD_REFUSALS = (
    TOO_MANY_ROWS,
    TOO_MANY_COLUMNS,
    TOO_MANY_CELLS,
    TOO_LONG_CELL,
    TOO_MUCH_TEXT,
    TOO_COSTLY,
)

# The child imports the project from the repository root, with its own limits set first. It sets
# its OOM score to the highest, so that if the container runs out of memory the child is the one
# the kernel kills, not the server.
_CHILD_PREAMBLE = (
    "import sys\n"
    "try:\n"
    "    import resource\n"
    "except ImportError:  # no memory cap on a platform without it (Windows): the time limit holds\n"
    "    resource = None\n"
    "cap = int(sys.argv[2])\n"
    "if resource is not None:\n"
    "    resource.setrlimit(resource.RLIMIT_AS, (cap, cap))\n"
    "try:\n"
    "    open('/proc/self/oom_score_adj', 'w').write('1000')\n"
    "except OSError:\n"
    "    pass\n"
)
_CHILD_CODE = _CHILD_PREAMBLE + (
    "sys.path.insert(0, sys.argv[3])\n"
    "from skills.excel.workbook import _scan_in_this_process\n"
    "sys.stdout.write(_scan_in_this_process(sys.argv[1], sys.argv[4]))\n"
)
# The limits the parent hands the child, by name: the child checks against the same numbers.
_LIMIT_NAMES = ("MAX_ROWS", "MAX_COLUMNS", "MAX_CELLS", "MAX_CELL_CHARS", "MAX_TEXT_BYTES")
_PROJECT_ROOT = str(Path(__file__).resolve().parents[2])


def _scan_in_child(input_path: str) -> str | None:
    """Run the readers on the workbook in a child process; the refusal text, or ``None``.

    The child runs ``_scan_in_this_process`` with an address-space cap (``SCAN_MEMORY_BYTES``) and
    is stopped after ``SCAN_SECONDS``. Out of memory, out of time, a crash, or an answer the
    parent doesn't know are all refusals: the file is never let through because its check failed.
    """
    environment = dict(os.environ)
    # One BLAS thread: each reserves address space the cap would count.
    environment.update(OPENBLAS_NUM_THREADS="1", OMP_NUM_THREADS="1", MKL_NUM_THREADS="1")
    try:
        done = subprocess.run(  # noqa: S603 — fixed argv, the path is an argument, not code
            [
                sys.executable,
                "-I",
                "-c",
                _CHILD_CODE,
                input_path,
                str(SCAN_MEMORY_BYTES),
                _PROJECT_ROOT,
                json.dumps({name: globals()[name] for name in _LIMIT_NAMES}),
            ],
            capture_output=True,
            text=True,
            timeout=SCAN_SECONDS,
            env=environment,
            check=False,
        )
    except subprocess.TimeoutExpired:
        logger.info("%r: the check ran past %d s", input_path, SCAN_SECONDS)
        return TOO_COSTLY
    except Exception as exc:  # noqa: BLE001 — the child could not be started: nothing was checked
        logger.exception("could not start the check of %r", input_path)
        # Refused, but not kept as the file's verdict: this says nothing about the file.
        raise WorkbookRefused(CHECK_STOPPED) from exc
    answer = done.stdout.strip()
    if done.returncode != 0:
        logger.info("%r: the check ended with status %d", input_path, done.returncode)
        return CHECK_STOPPED
    if answer == "PASS":
        return None
    prefix, _, index = answer.partition(":")
    if prefix == "REFUSE" and index.isdigit() and int(index) < len(_CHILD_REFUSALS):
        return _CHILD_REFUSALS[int(index)]
    return CHECK_STOPPED


def _scan_in_this_process(input_path: str, limits: str = "{}") -> str:
    """The child's job: ``"PASS"`` or ``"REFUSE:<n>"`` (``n`` indexes ``_CHILD_REFUSALS``).

    ``limits`` is the JSON the parent passes: the values of ``_LIMIT_NAMES`` to check against.

    Counts what the table's reader would build (``_reader_verdict``), then reads the workbook
    the ways the skills, the verifiers and the analysis do (``_read_as_the_skills_do``), so what
    is held in memory here is what the server's own reads will hold. Only an out-of-memory error
    is an answer of its own; a file the readers can't read is left to the callers' refusals.
    """
    global _in_scan_child
    _in_scan_child = True
    for name, value in json.loads(limits).items():
        if name in _LIMIT_NAMES and isinstance(value, int):
            globals()[name] = value
    try:
        verdict = _reader_verdict(input_path)
        if verdict is None:
            _read_as_the_skills_do(input_path)
    except MemoryError:
        verdict = TOO_COSTLY
    if verdict is None:
        return "PASS"
    return f"REFUSE:{_CHILD_REFUSALS.index(verdict)}"


def _read_as_the_skills_do(input_path: str) -> None:
    """Build the tables the analysis builds, and the sheet listing the planner reads.

    The analysis holds a default read and a ``dtype=object`` read of the first sheet at once
    (``SpreadsheetAnalysisSkill``, the largest the app makes); the cleaning skills and the
    Inspector make the second alone; the planner and the Inspector list the sheets. The layout
    check and the formula count only stream cells one row at a time over the same load, so they
    add time, not memory, and are not repeated here. A file these cannot read is not this
    check's business, so every error but running out of memory is left to the callers.
    """
    import pandas as pd

    try:
        sheet_names(input_path)
        frames = (pd.read_excel(input_path), pd.read_excel(input_path, dtype=object))
        del frames
    except MemoryError:
        raise
    except Exception:  # noqa: BLE001 — unreadable: the callers' own reads refuse it, as they do
        logger.info("%r couldn't be read by the check", input_path)


class _LimitPassed(Exception):
    """Raised inside a scan to stop at the first limit passed."""


def _reader_verdict(input_path: str) -> str | None:
    """Walk the workbook the way the table is read, counting what that read would build.

    pandas builds its table from openpyxl's read-only rows, so this uses the same reader on the
    same file and counts what it yields, instead of guessing at the XML: the rows it pads out to
    (one cell at row 20,000 is 19,999 rows), the widest row, and the text in the cells. It holds one
    row at a time and stops at the first limit passed. A workbook the reader can't read is left
    to the callers' own refusals, because pandas can't read it either.
    """
    try:
        with open(input_path, "rb") as handle:
            book = openpyxl.load_workbook(
                handle, read_only=True, data_only=True, keep_links=False
            )
            try:
                grids = chars = all_rows = 0
                for sheet in book.worksheets:
                    sheet.reset_dimensions()  # as pandas does: a wrong <dimension> hides nothing
                    rows = widest = 0
                    for row in sheet.iter_rows(values_only=True):
                        rows += 1
                        all_rows += 1
                        widest = max(widest, len(row))
                        # All sheets together: one part can be declared as thousands of sheets,
                        # and the reader pads every one of them out to its last row.
                        if all_rows > MAX_ROWS:
                            return TOO_MANY_ROWS
                        if widest > MAX_COLUMNS:
                            return TOO_MANY_COLUMNS
                        if grids + rows * widest > MAX_CELLS:
                            return TOO_MANY_CELLS
                        for value in row:
                            if isinstance(value, str):
                                if len(value) > MAX_CELL_CHARS:
                                    return TOO_LONG_CELL
                                chars += len(value) * _bytes_a_character(value)
                                if chars > MAX_TEXT_BYTES:
                                    return TOO_MUCH_TEXT
                    grids += rows * widest
            finally:
                book.close()
    except MemoryError:
        raise  # not "can't be read": the file needs more than the check was given
    except Exception:  # noqa: BLE001 — the same reader would fail in pandas, with a refusal
        logger.info("%r couldn't be walked for size limits", input_path)
    return None


def _bytes_a_character(text: str) -> int:
    """How many bytes CPython stores for each character of ``text``: 1, 2 or 4."""
    widest = max(text, default="")
    return 4 if widest > "\uffff" else 2 if widest > "\xff" else 1


def _zip_sheet_names(handle) -> list[str]:  # noqa: ANN001 — a binary file object
    """The declared sheet names of an .xlsx, cross-checked with openpyxl."""
    with zipfile.ZipFile(handle) as archive:
        declared = _declared_sheet_names(archive)
    if not declared:
        raise WorkbookRefused(_NO_SHEET)
    handle.seek(0)
    # A file handle, not the path: openpyxl rejects an .xls name even when the
    # content is a workbook it can read.
    book = openpyxl.load_workbook(handle, read_only=True, keep_links=False)
    try:
        loaded = list(book.sheetnames)
    finally:
        book.close()
    if loaded != declared:
        logger.info("declared sheets %r, openpyxl loaded %r", declared, loaded)
        raise WorkbookRefused(_SHEETS_UNREADABLE)
    return declared


def _declared_sheet_names(archive: zipfile.ZipFile) -> list[str]:
    """The ``<sheet name=…>`` entries of the workbook part, in order."""
    content_types = ElementTree.fromstring(_read_part(archive, "[Content_Types].xml"))
    parts = [
        element.get("PartName", "")
        for element in content_types.iter()
        if _local(element.tag) == "Override"
        and element.get("ContentType") in _WORKBOOK_CONTENT_TYPES
    ]
    if len(parts) != 1:
        raise WorkbookRefused(_UNREADABLE)
    workbook = ElementTree.fromstring(_read_part(archive, parts[0].lstrip("/")))
    return [
        sheet.get("name", "")
        for sheets in workbook.iter()
        if _local(sheets.tag) == "sheets"
        for sheet in sheets
        if _local(sheet.tag) == "sheet"
    ]


def _read_part(archive: zipfile.ZipFile, name: str) -> bytes:
    """One part of the archive, refusing an absurdly large one."""
    with archive.open(name) as part:
        data = part.read(_MAX_PART_BYTES + 1)
    if len(data) > _MAX_PART_BYTES:
        raise WorkbookRefused(_UNREADABLE)
    return data


def _local(tag: str) -> str:
    """An XML tag without its ``{namespace}`` prefix."""
    return tag.rpartition("}")[2]


def _has_stream(handle) -> bool:  # noqa: ANN001 — a binary file object
    """Whether the OLE2 file holds the ``EncryptedPackage`` stream's name."""
    overlap = len(_ENCRYPTED_STREAM) - 1
    tail = b""
    while chunk := handle.read(1 << 20):
        if _ENCRYPTED_STREAM in tail + chunk:
            return True
        tail = chunk[-overlap:]
    return False


def _sheet_problem(input_path: str, sheet_name: str) -> str | None:
    """The refusal text for a sheet that isn't a plain table, or ``None``."""
    try:
        with open(input_path, "rb") as handle:
            book = openpyxl.load_workbook(
                handle, read_only=True, data_only=True, keep_links=False
            )
            try:
                sheet = book[sheet_name]
                # A wrong <dimension> would hide cells that pandas (which ignores it) reads.
                sheet.reset_dimensions()
                reason = _MERGED_CELLS if _has_merged_cells(sheet) else _layout_problem(sheet)
            finally:
                book.close()
    except MemoryError:
        raise
    except Exception as exc:  # noqa: BLE001 — unreadable or corrupt, reported as a refusal
        logger.exception("could not check the layout of %r", input_path)
        raise WorkbookRefused(_UNREADABLE) from exc
    if reason is None or reason == _READS_AS_EMPTY:
        return reason
    return reason + _PLAIN_TABLES_ONLY


def _is_empty(value: object) -> bool:
    """Whether pandas reads this cell as missing: nothing, or an empty string."""
    return value is None or value == ""


def _layout_problem(sheet) -> str | None:  # noqa: ANN001 — an openpyxl read-only sheet
    """Which plain-table rule the sheet's cells break (merged cells aside), if any."""
    headers: dict[int, object] = {}  # column -> its value on row 1
    data_columns: set[int] = set()  # columns with a value below row 1
    seen_first_row = False
    for row in sheet.iter_rows():
        for cell in row:
            if _is_empty(cell.value):
                continue
            if cell.data_type == "e" or (
                cell.row != 1 and isinstance(cell.value, str) and cell.value in _PANDAS_NA_TEXT
            ):
                return _READS_AS_EMPTY
            if not seen_first_row:
                if cell.row != 1:
                    return _HEADER_NOT_FIRST  # nothing on row 1, something below it
                seen_first_row = True
            if cell.row == 1:
                headers[cell.column] = cell.value
            else:
                data_columns.add(cell.column)
    if not headers:
        return None  # an empty sheet: nothing to change
    last_header = max(headers)
    if len(headers) != last_header or any(column > last_header for column in data_columns):
        return _COLUMN_WITHOUT_HEADER
    if len(set(headers.values())) != len(headers):
        return _DUPLICATE_HEADER
    return None


class _MergedCellFound(Exception):
    """Raised inside the XML scan to stop at the first ``<mergeCell>``."""


def _has_merged_cells(sheet) -> bool:  # noqa: ANN001 — an openpyxl read-only sheet
    """Whether the sheet's XML declares a merged range.

    openpyxl's read-only mode doesn't expose merged cells, so the sheet's own
    part (opened through openpyxl) is scanned for ``<mergeCell>`` with a
    streaming parser that builds no tree, so a big sheet costs no memory.
    """
    parser = expat.ParserCreate(namespace_separator=" ")

    def on_start(name: str, attributes: dict[str, str]) -> None:
        if name.rpartition(" ")[2] == "mergeCell":
            raise _MergedCellFound

    parser.StartElementHandler = on_start
    with sheet._get_source() as source:
        try:
            while chunk := source.read(_SCAN_CHUNK_BYTES):
                parser.Parse(chunk, False)
            parser.Parse(b"", True)
        except _MergedCellFound:
            return True
    return False


@dataclass(frozen=True)
class FormulaCounts:
    """The formula cells of a workbook, as two fresh reads of it count them.

    Attributes:
        formulas: Cells that hold a formula.
        without_value: Of those, the cells with no saved result (never
            calculated, or an empty one): a values-only copy has an empty
            cell there.
    """

    formulas: int
    without_value: int


def formula_cell_counts(input_path: str) -> FormulaCounts:
    """Count the formula cells of every sheet, and those with no saved value.

    Two read-only openpyxl passes over the same file, in step: one showing
    formulas, one showing the values the file saved for them (what a pandas
    read, and so a values-only copy, gets).

    Args:
        input_path: The workbook to read. It is only ever read.

    Returns:
        The counts.

    Raises:
        WorkbookRefused: The workbook can't be read (see ``sheet_names``).
    """
    names = sheet_names(input_path)
    formulas = without_value = 0
    try:
        with open(input_path, "rb") as shown_handle, open(input_path, "rb") as saved_handle:
            shown_book = openpyxl.load_workbook(
                shown_handle, read_only=True, data_only=False, keep_links=False
            )
            saved_book = openpyxl.load_workbook(
                saved_handle, read_only=True, data_only=True, keep_links=False
            )
            try:
                for name in names:
                    shown_sheet, saved_sheet = shown_book[name], saved_book[name]
                    # A wrong <dimension> would hide cells that pandas (which ignores it) reads.
                    shown_sheet.reset_dimensions()
                    saved_sheet.reset_dimensions()
                    for shown_row, saved_row in zip(
                        shown_sheet.iter_rows(), saved_sheet.iter_rows(), strict=True
                    ):
                        for shown, saved in zip(shown_row, saved_row, strict=True):
                            if shown.data_type == "f":
                                formulas += 1
                                without_value += _is_empty(saved.value)
            finally:
                shown_book.close()
                saved_book.close()
    except MemoryError:
        raise
    except Exception as exc:  # noqa: BLE001 — unreadable or corrupt, reported as a refusal
        logger.exception("could not count the formulas of %r", input_path)
        raise WorkbookRefused(_UNREADABLE) from exc
    return FormulaCounts(formulas=formulas, without_value=without_value)


def discard_partial_output(output_path: Path) -> None:
    """Delete what a failed write left at ``output_path``, if anything.

    The path is always a freshly built one under ``workspace/outputs/``
    (never the upload), and a failed write must not leave a half-written
    file there that could be mistaken for a result.
    """
    try:
        output_path.unlink(missing_ok=True)
    except OSError:
        logger.exception("could not remove the partial output %r", output_path)
