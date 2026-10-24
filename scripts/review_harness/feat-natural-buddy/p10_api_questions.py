"""P10 (round 3): Brain's new `question_about_file`, through the HTTP API (real app, real pipeline,
fake models). The backend offers the latest copy (FileChoice), which the engine-level harness
(`nb.py`) doesn't, so the question turns are probed here.

    python p10_api_questions.py [--tree ../../..] [--out out/r3-p10.json] [CASE ...]

Each case: turns of (message, Brain/Think replies), a writer (careful echo, or echo plus social
words), and the reviewer's truth label for the LAST turn's reply. Records the messages, the
reply source, delivered files, the uploaded original's SHA-256 before/after, and the latest
delivered copy read with openpyxl (rows, duplicates, blank rows).
"""
import argparse
import hashlib
import json
import sys
import tempfile
from pathlib import Path

ap = argparse.ArgumentParser()
ap.add_argument("--tree", default=str(Path(__file__).resolve().parents[3]))
ap.add_argument("--out")
ap.add_argument("--modes", default="friend,focus")
ap.add_argument("--question-intent", default="question_about_file",
                help="Brain's reading of the question (base has no question_about_file: try conversation, unclear)")
ap.add_argument("only", nargs="*")
a = ap.parse_args()
sys.path.insert(0, str(Path(a.tree).resolve()))

from fastapi.testclient import TestClient  # noqa: E402

from backend.main import create_app  # noqa: E402
from backend.settings import Settings  # noqa: E402
from tests.backend.factories import DEDUPE_THEN_EMPTY, PipelineFactory, QueuedProvider, chain, single  # noqa: E402
from tests.backend.helpers import RELAXED_LIMITS, new_session, start_turn, upload, wait_finished, xlsx_bytes  # noqa: E402

try:
    from tests.buddy.writer_fakes import echo_answer, facts_in  # noqa: E402
except ImportError:  # base: no writer
    echo_answer = facts_in = None


def social(prefix="", suffix="", after_kind=None):
    """Echo the required facts by slot; add social words before the first and after the last
    (or after the fact of `after_kind`)."""
    def w(prompt):
        facts = facts_in(prompt)
        if not facts:
            return json.dumps({"messages": [(prefix + " hey").strip()]})
        req = [(f[0], f[2], f"{{{f[3][0]}}}") for f in facts if f[1]]
        msgs = [s for _, _, s in req]
        if len(msgs) > 3:
            msgs = msgs[:2] + [" ".join(msgs[2:])]
        if prefix:
            msgs[0] = f"{prefix} {msgs[0]}"
        if suffix:
            idx = len(msgs) - 1
            if after_kind:
                for i, (fid, kind, s) in enumerate(req):
                    if kind.startswith(after_kind):
                        idx = min(i, len(msgs) - 1)
            msgs[idx] = f"{msgs[idx]} {suffix}"
        return json.dumps({"messages": msgs})
    return w


class W:
    def __init__(self, fn):
        self.fn, self.calls = fn, 0

    def complete(self, *, system, messages, max_tokens=1024):
        self.calls += 1
        return self.fn(messages[0].content)


DEDUPE = ("remove duplicate rows", single("remove_duplicate_rows"))
BOTH = ("remove duplicate rows and then remove empty rows", chain(*DEDUPE_THEN_EMPTY))
Q = lambda m: (m, single(a.question_intent))  # noqa: E731

CASES = [
    # (id, turns, writer fn, truth for the last reply, what makes it so)
    ("Q1-them-echo", [DEDUPE, Q("did you get rid of them?")], "echo", "TRUE", ""),
    ("Q2-dupes-echo", [DEDUPE, Q("did you get rid of the duplicates?")], "echo", "?", "says 'I didn't run the duplicate-row removal' after a dedupe"),
    ("Q3-empty-echo", [DEDUPE, Q("did you also remove the empty row?")], "echo", "TRUE", "copy keeps 1 empty row"),
    ("Q4-clean-echo", [BOTH, Q("is the file clean now?")], "echo", "TRUE", ""),
    ("Q5-first-echo", [Q("did you get rid of the duplicates?")], "echo", "TRUE", "nothing done yet"),
    ("Q6-empty-ya", [DEDUPE, Q("did you also remove the empty row?")], ("ya totally!", ""), "FALSE", "'yes' to a part not done"),
    ("Q7-empty-ididit", [DEDUPE, Q("did you also remove the empty row?")], ("i did!", ""), "FALSE", "'yes' to a part not done"),
    ("Q8-empty-bilkul", [DEDUPE, Q("did you also remove the empty row?")], ("bilkul bhai!", ""), "FALSE", "'yes' (Hinglish)"),
    ("Q9-first-true", [Q("did you get rid of the duplicates?")], ("true, i did!", ""), "FALSE", "nothing was done"),
    ("Q10-them-no", [DEDUPE, Q("did you get rid of them?")], ("no, not yet 😅", ""), "FALSE", "they were removed"),
    ("Q11-them-never", [DEDUPE, Q("did you get rid of them?")], ("nah i never did", ""), "FALSE", "they were removed ('nah' may not pass)"),
    ("Q12-first-went-away", [Q("did you get rid of the duplicates?")], ("", "they all went away though 👍"), "FALSE", "nothing was done"),
    ("Q13-control", [DEDUPE, Q("did you also remove the empty row?")], ("not yet!", "want me to?"), "TRUE", ""),
]


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def read(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows = [tuple(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)][1:]
    wb.close()
    seen, d = set(), 0
    for r in rows:
        if r in seen and any(c is not None for c in r):
            d += 1
        seen.add(r)
    return {"rows": len(rows), "dups": d, "blank": sum(all(c is None for c in r) for r in rows)}


res = []
for cid, turns, wspec, truth, why in CASES:
    if a.only and not any(cid.startswith(o) for o in a.only):
        continue
    for mode in a.modes.split(","):
        tmp = Path(tempfile.mkdtemp())
        settings = Settings(data_dir=tmp / "data", rate_limits=RELAXED_LIMITS)
        replies = [r for _, rs in turns for r in rs]
        if echo_answer is None:
            factory = PipelineFactory(QueuedProvider(replies), speaker=mode)
            writer = None
        else:
            fn = echo_answer if wspec == "echo" else social(*wspec)
            writer = W(fn)
            factory = PipelineFactory(QueuedProvider(replies), speaker="natural", writer=writer)
        client = TestClient(create_app(settings, factory))
        sid = new_session(client, mode)
        assert upload(client, sid, xlsx_bytes()).status_code == 201
        originals = [p for p in (tmp / "data").rglob("*.xlsx")]
        h0 = {str(p): sha(p) for p in originals}
        out = []
        for msg, _ in turns:
            st = start_turn(client, sid, msg)
            out.append(wait_finished(client, sid, st.json()["turn_id"]))
        last = out[-1]
        delivered = [f for t in out for f in (t.get("delivered_files") or [])]
        copies = sorted((p for p in (tmp / "data").rglob("*.xlsx") if str(p) not in h0), key=lambda p: p.stat().st_mtime)
        rec = dict(case=cid, mode=mode, truth=truth, why=why, outcome=last["outcome"],
                   source=last.get("reply_source"), messages=last.get("messages") or [last["reply"]],
                   delivered_last=len(last.get("delivered_files") or []),
                   latest_copy=read(copies[-1]) if copies else None,
                   orig_unchanged=all(sha(p) == h for p, h in h0.items()),
                   writer_calls=writer.calls if writer else 0)
        res.append(rec)
        print(f"{cid:22} {mode:6} {truth:5} src={rec['source']} out={rec['outcome']} calls={rec['writer_calls']} "
              f"copy={rec['latest_copy']} orig_ok={rec['orig_unchanged']}")
        for m in rec["messages"]:
            print("      |", m)
if a.out:
    Path(a.out).parent.mkdir(exist_ok=True)
    json.dump(res, open(a.out, "w"), indent=1, ensure_ascii=False)
print("original changed:", [r["case"] for r in res if not r["orig_unchanged"]] or "none")
