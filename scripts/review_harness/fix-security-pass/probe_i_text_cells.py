"""Text that looks like a number, through both cleaning skills: what do the copy's cells hold?

    python scripts/review_harness/fix-security-pass/probe_i_text_cells.py [TREE]
"""
import logging
import os
import sys
import tempfile
from pathlib import Path

tree = str(Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[3]).resolve())
sys.path.insert(0, tree)
logging.disable(logging.CRITICAL)
import openpyxl  # noqa: E402

from skills.base import PermissionTier, SkillTask  # noqa: E402
from skills.excel.dedupe import ExcelDeduplicateSkill  # noqa: E402
from skills.excel.remove_empty_rows import ExcelRemoveEmptyRowsSkill  # noqa: E402

d = Path(tempfile.mkdtemp())
os.chdir(d)
src = d / "in.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ROWS = [("code", "n", "when"), ("007", 1, "2024-01-02"), (7, 1, "2024-01-02"), ("A", 2, None), ("A", 2, None),
        ("123456789012345678", 3, None), (123456789012345678, 3, None), (None, None, None), ("1e3", 1000, None),
        ("TRUE", True, None), (True, True, None), ("0.10", 0.1, None), (0.1, 0.1, None)]
for r in ROWS:
    ws.append(r)
wb.save(src)
for skill in (ExcelDeduplicateSkill(), ExcelRemoveEmptyRowsSkill()):
    task = SkillTask(task_id="t-1", session_id="s", workspace_id="w", description="x",
                     input_references=[str(src)], permission=PermissionTier.WRITE, output_dir=str(d))
    res = skill.execute(task)
    out = res.outputs[0] if res.outputs else None
    print(f"== {type(skill).__name__}: status={res.status} summary={getattr(res, 'summary', '')!r}")
    if out:
        rows = list(openpyxl.load_workbook(out).active.iter_rows(values_only=True))
        for r in rows:
            print("   ", tuple(f"{v!r}:{type(v).__name__}" for v in r))
