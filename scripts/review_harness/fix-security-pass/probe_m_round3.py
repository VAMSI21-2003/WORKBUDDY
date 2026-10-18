"""Round 3: workbooks aimed at the reader-based check (`_reader_verdict`, head `c766fa6`).

For each file: size, does backend.uploads.check_workbook (the upload route's structural check) accept it,
how long check_workbook_limits takes and what it says, and its peak RSS. Each file is checked in a fresh
child process (so the cache and RSS are per file), capped at --cap-mib of address space and --timeout
seconds. Run from the repo root with the image's parser:

    OPENPYXL_LXML=False python scripts/review_harness/fix-security-pass/probe_m_round3.py [--tree DIR] [--only M1,M3] [--read]

--read also times one pandas read (as the skills do it) of every file the check passes.
Files are written to out/m_*.xlsx (probe_l_turns.py-style turns: probe_n_turns.py).
"""
import argparse
import os
import subprocess
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "out"
OUT.mkdir(exist_ok=True)

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"


def header_row(cols: int) -> str:
    from openpyxl.utils import get_column_letter
    return "".join(f'<c r="{get_column_letter(i)}1" t="inlineStr"><is><t>h{i}</t></is></c>'
                   for i in range(1, cols + 1))


def bomb_sheet(rows: int = 20_000, cols: int = 1000) -> str:
    """1,000 headers in row 1 and one cell at row ``rows``: refused by the walk's grid."""
    return (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{NS}"><sheetData>'
            f'<row r="1">{header_row(cols)}</row><row r="{rows}"><c><v>1</v></c></row>'
            "</sheetData></worksheet>")


def table_sheet(rows: list[list[object]]) -> str:
    from openpyxl.utils import get_column_letter
    from xml.sax.saxutils import escape
    out = []
    for r, values in enumerate(rows, 1):
        cells = []
        for c, v in enumerate(values, 1):
            ref = f"{get_column_letter(c)}{r}"
            if isinstance(v, str):
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(v)}</t></is></c>')
            else:
                cells.append(f'<c r="{ref}"><v>{v}</v></c>')
        out.append(f'<row r="{r}">{"".join(cells)}</row>')
    return (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{NS}"><sheetData>'
            f'{"".join(out)}</sheetData></worksheet>')


def package(path: Path, sheet: str, *, sheets: int = 1, wb_extra: str = "", rels_extra: str = "",
            ct_extra: str = "", parts: dict[str, bytes] | None = None) -> None:
    """A minimal workbook. ``sheets`` > 1 declares that many sheets, all on the one sheet part."""
    sheet_entries = "".join(f'<sheet name="S{i}" sheetId="{i}" r:id="rId{i}"/>' for i in range(1, sheets + 1))
    rels = "".join(f'<Relationship Id="rId{i}" Type="{R}/worksheet" Target="worksheets/sheet1.xml"/>'
                   for i in range(1, sheets + 1))
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.sheet.main+xml"/>'
            f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{CT_SHEET}"/>'
            f"{ct_extra}</Types>").encode(),
        "_rels/.rels": (f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PR}">'
                        f'<Relationship Id="rId1" Type="{R}/officeDocument" Target="xl/workbook.xml"/>'
                        "</Relationships>").encode(),
        "xl/workbook.xml": (f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="{NS}" xmlns:r="{R}">'
                            f"<sheets>{sheet_entries}</sheets>{wb_extra}</workbook>").encode(),
        "xl/_rels/workbook.xml.rels": (f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PR}">'
                                       f"{rels}{rels_extra}</Relationships>").encode(),
        "xl/worksheets/sheet1.xml": sheet.encode(),
        **(parts or {}),
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for name, data in files.items():
            z.writestr(name, data)


EMPTY_FAR_ROW = (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{NS}"><sheetData>'
                 '<row r="200000"/></sheetData></worksheet>')


def m1_many_sheets(path: Path, n: int) -> None:
    """``n`` sheets declared on ONE sheet part (so the upload route sees one worksheet), each part
    a single empty <row r="200000"/>: every sheet walks 200,000 empty rows, grid 0, no limit hit."""
    package(path, EMPTY_FAR_ROW, sheets=n)


EXT_REL = f'<Relationship Id="rIdX" Type="{R}/externalLink" Target="externalLinks/externalLink1.xml"/>'
EXT_WB = '<externalReferences><externalReference r:id="rIdX"/></externalReferences>'


def m2_broken_link_bomb(path: Path) -> None:
    """The grid bomb, plus an external reference whose part is missing. The walk (keep_links=True,
    openpyxl's default) raises and returns no verdict; pandas (keep_links=False) never reads links."""
    package(path, bomb_sheet(), wb_extra=EXT_WB, rels_extra=EXT_REL)


def m3_many_shared_strings(path: Path, n: int) -> None:
    """``n`` distinct shared strings, none referenced except the first: openpyxl loads every one."""
    sst = (f'<sst xmlns="{NS}">' + "<si><t>ab</t></si>" * n + "</sst>").encode()
    sheet = table_sheet([["name"], ["a"]]).replace('t="inlineStr"><is><t>a</t></is>', 't="s"><v>0</v>')
    package(path, sheet, rels_extra=f'<Relationship Id="rIdS" Type="{R}/sharedStrings" Target="sharedStrings.xml"/>',
            ct_extra='<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                     'officedocument.spreadsheetml.sharedStrings+xml"/>',
            parts={"xl/sharedStrings.xml": sst})


def m4_many_styles(path: Path, n: int) -> None:
    """``n`` cell formats in styles.xml: openpyxl builds an object for every one on load."""
    xf = '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    styles = (f'<styleSheet xmlns="{NS}"><fonts count="1"><font/></fonts><fills count="1"><fill/></fills>'
              '<borders count="1"><border/></borders><cellStyleXfs count="1"><xf/></cellStyleXfs>'
              f'<cellXfs count="{n}">' + xf * n + "</cellXfs></styleSheet>").encode()
    package(path, table_sheet([["name"], ["a"]]),
            rels_extra=f'<Relationship Id="rIdY" Type="{R}/styles" Target="styles.xml"/>',
            ct_extra='<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-'
                     'officedocument.spreadsheetml.styles+xml"/>',
            parts={"xl/styles.xml": styles})


def m5_big_external_link(path: Path, n: int) -> None:
    """An external link whose cached sheet holds ``n`` cells: read whole by openpyxl's default load
    (the walk and sheet_names), never by pandas."""
    cells = '<cell r="A1"><v>1</v></cell>' * n
    link = (f'<externalLink xmlns="{NS}" xmlns:r="{R}"><externalBook r:id="rId1"><sheetNames>'
            f'<sheetName val="S"/></sheetNames><sheetDataSet><sheetData sheetId="0"><row r="1">{cells}</row>'
            "</sheetData></sheetDataSet></externalBook></externalLink>").encode()
    link_rels = (f'<Relationships xmlns="{PR}"><Relationship Id="rId1" Type="{R}/externalLinkPath" '
                 'Target="other.xlsx" TargetMode="External"/></Relationships>').encode()
    package(path, table_sheet([["name"], ["a"]]), wb_extra=EXT_WB, rels_extra=EXT_REL,
            parts={"xl/externalLinks/externalLink1.xml": link,
                   "xl/externalLinks/_rels/externalLink1.xml.rels": link_rels})


def m6_long_names(path: Path, rows: int, chars: int) -> None:
    """name/email/phone table whose names are ``chars`` long and share their first 3 letters, with
    near-duplicate pairs: the near-duplicate check compares long strings pair by pair."""
    data = [["name", "email", "phone"]]
    for i in range(rows):
        data.append(["abc" + "x" * (chars - 10) + f"{i // 2:07d}", f"u{i // 2}@x.com", f"555{i // 2:04d}"])
    package(path, table_sheet(data))


def m7_text_at_cap(path: Path) -> None:
    """1,500 cells of 32,767 characters in one column (49 million characters: inside both caps),
    as shared strings so the file stays small; plus an id column."""
    n = 1500
    sst = (f'<sst xmlns="{NS}"><si><t>text</t></si><si><t>id</t></si>'
           + "".join(f"<si><t>{'y' * 32_760}{i:07d}</t></si>" for i in range(n)) + "</sst>").encode()
    rows = ['<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>']
    rows += [f'<row r="{r}"><c r="A{r}" t="s"><v>{r}</v></c><c r="B{r}"><v>{r % 700}</v></c></row>'
             for r in range(2, n + 2)]
    sheet = (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{NS}"><sheetData>{"".join(rows)}'
             "</sheetData></worksheet>")
    package(path, sheet, rels_extra=f'<Relationship Id="rIdS" Type="{R}/sharedStrings" Target="sharedStrings.xml"/>',
            ct_extra='<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                     'officedocument.spreadsheetml.sharedStrings+xml"/>',
            parts={"xl/sharedStrings.xml": sst})


def m8_defined_names(path: Path, n: int) -> None:
    """``n`` defined names in workbook.xml: parsed whole on every load."""
    names = '<definedName name="n">S1!$A$1</definedName>' * n
    package(path, table_sheet([["name"], ["a"]]), wb_extra=f"<definedNames>{names}</definedNames>")


def m0_ordinary(path: Path) -> None:
    package(path, table_sheet([["name", "amount"]] + [[f"n{i % 40}", i % 40] for i in range(50)]))


CASES = {
    "M0_ordinary_control": m0_ordinary,
    "M1_sheets_100": lambda p: m1_many_sheets(p, 100),
    "M1_sheets_2000": lambda p: m1_many_sheets(p, 2000),
    "M1_sheets_20000": lambda p: m1_many_sheets(p, 20000),
    "M2_broken_link_bomb": m2_broken_link_bomb,
    "M3_shared_strings_3_6m": lambda p: m3_many_shared_strings(p, 3_600_000),
    "M4_styles_1m": lambda p: m4_many_styles(p, 1_050_000),
    "M5_ext_link_2_3m_cells": lambda p: m5_big_external_link(p, 2_300_000),
    "M6_long_names_200x30k": lambda p: m6_long_names(p, 200, 30_000),
    "M7_text_at_cap": m7_text_at_cap,
    "M8_defined_names_1_4m": lambda p: m8_defined_names(p, 1_400_000),
}

CHECK = r"""
import resource, sys, time
cap = int(sys.argv[3]) * 1024 * 1024
resource.setrlimit(resource.RLIMIT_AS, (cap, cap))
sys.path.insert(0, sys.argv[2])
from pathlib import Path
from backend.settings import Settings
from backend.uploads import check_workbook
try:
    import tempfile
    check_workbook(Path(sys.argv[1]), Settings(data_dir=Path(tempfile.mkdtemp())))
    upload = "accepted"
except Exception as e:
    upload = f"rejected({getattr(e, 'code', type(e).__name__)})"
size = Path(sys.argv[1]).stat().st_size
if size > Settings.__dataclass_fields__["max_upload_bytes"].default:
    upload += "(over-upload-cap)"
from skills.excel.workbook import check_workbook_limits
t = time.time()
try:
    check_workbook_limits(sys.argv[1]); verdict = "passes"
except MemoryError:
    verdict = "MemoryError"
except Exception as e:
    verdict = f"refused: {e}"
print(f"upload={upload} check {time.time() - t:.1f}s peak={resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024:.0f}MiB {verdict}")
"""

READ = r"""
import resource, sys, time
cap = int(sys.argv[2]) * 1024 * 1024
resource.setrlimit(resource.RLIMIT_AS, (cap, cap))
import pandas as pd
t = time.time()
try:
    shape = pd.read_excel(sys.argv[1], dtype=object).shape
except MemoryError:
    shape = "MemoryError"
except Exception as e:
    shape = type(e).__name__
print(f"pandas {shape} {time.time() - t:.1f}s peak={resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024:.0f}MiB")
"""


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tree", default=str(HERE.parents[2]))
    ap.add_argument("--read", action="store_true")
    ap.add_argument("--cap-mib", type=int, default=3072)
    ap.add_argument("--timeout", type=int, default=600)
    ap.add_argument("--only", default="")
    a = ap.parse_args()
    only = set(filter(None, a.only.split(",")))
    env = dict(os.environ, OPENPYXL_LXML="False")
    for name, make in CASES.items():
        if only and name.split("_")[0] not in only and name not in only:
            continue
        path = OUT / f"m_{name}.xlsx"
        if not path.exists():
            make(path)
        with zipfile.ZipFile(path) as z:
            unpacked = sum(i.file_size for i in z.infolist())
        line = f"{name:26s} {path.stat().st_size:9d} B zipped, {unpacked / 2**20:6.1f} MiB unpacked | "
        try:
            r = subprocess.run([sys.executable, "-c", CHECK, str(path), a.tree, str(a.cap_mib)],
                               capture_output=True, text=True, env=env, timeout=a.timeout)
            line += r.stdout.strip() or ("ERR " + r.stderr.strip()[-300:])
        except subprocess.TimeoutExpired:
            line += f"check still running at {a.timeout}s (killed)"
        if a.read and "passes" in line:
            try:
                r = subprocess.run([sys.executable, "-c", READ, str(path), str(a.cap_mib)],
                                   capture_output=True, text=True, env=env, timeout=a.timeout)
                line += " | " + (r.stdout.strip() or r.stderr.strip()[-200:])
            except subprocess.TimeoutExpired:
                line += f" | pandas still running at {a.timeout}s"
        print(line, flush=True)


if __name__ == "__main__":
    main()
