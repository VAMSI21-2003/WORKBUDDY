"""Round 1 adversarial probes for fix/values-only-copy (needs values_only_probe.py next to it).

  python scripts/review_harness/adversarial_probe.py [NAME ...]

Cases: formula-count read failing (with a secret in the exception text), skills that don't refuse
(Inspector independence, single step and chain), shared formulas, #N/A saved results, formulas that
return "", an XML entity bomb in the sheet part, and timing on a large sheet.
"""
import re
import sys
import time
import zipfile

_args = sys.argv[1:]
_fwd = []
if "--tree" in _args:  # --tree DIR is passed on to values_only_probe / harness
    _i = _args.index("--tree")
    _fwd = _args[_i:_i + 2]
    del _args[_i:_i + 2]
ONLY = [a for a in _args if not a.startswith("--")]
sys.argv = [sys.argv[0]] + _fwd
import values_only_probe as V  # noqa: E402  (same directory)

H = V.H
FX = V.FX


def show(title, rec):
    print(f"\n## {title}\n  delivered={rec.get('delivered')} steps={rec.get('steps')} orig_unchanged={rec.get('orig_unchanged')} err={rec.get('error')}")
    print("  reply:", rec.get("reply"))
    if rec.get("delivered"):
        print("  out:", rec.get("out_cells"))


def case_measure_fails():
    import buddy.inspector as I

    def boom(path):
        raise RuntimeError("sk-SECRET-TOKEN-123 at /etc/passwd")

    real = I.formula_cell_counts
    I.formula_cell_counts = boom
    try:
        for sp in ("template", "friend", "focus"):
            show(f"formula count raises ({sp})", V.run_one("formulas_saved", "dedupe", sp, "s1"))
    finally:
        I.formula_cell_counts = real


def case_skill_doesnt_refuse():
    import skills.excel.dedupe as D
    import skills.excel.remove_empty_rows as R
    from skills.excel.workbook import single_sheet_name
    real = (D.plain_table_sheet_name, R.plain_table_sheet_name)
    D.plain_table_sheet_name = single_sheet_name
    R.plain_table_sheet_name = single_sheet_name
    try:
        for fx in ("merged", "title_row", "dup_header", "header_gap"):
            for req in ("dedupe", "chain"):
                show(f"skills don't refuse: {fx}/{req}", V.run_one(fx, req, "template", "s1"))
        show("skills don't refuse: merged/dedupe (focus)", V.run_one("merged", "dedupe", "focus", "s1"))
    finally:
        D.plain_table_sheet_name, R.plain_table_sheet_name = real


@V.fx("shared_formulas")
def _(p):
    wb, ws = V.base_book(("name", "amount", "double"), rows=[("Alice", 10, None), ("Bob", 20, None), ("Alice", 10, None), ("Carol", 30, None)])
    wb.save(p)

    def f(x):
        x = re.sub(r'<c r="C2"[^>]*/>|<c r="C2"[^>]*></c>', "", x)
        return x.replace('<c r="A2"', '<c r="C2"><f t="shared" ref="C2:C5" si="0">B2*2</f><v>20</v></c><c r="A2"', 1) if False else x
    # build XML by hand: replace rows 2..5 col C
    def g(x):
        for i, v in zip(range(2, 6), (20, 40, 20, 60)):
            fx = f'<f t="shared" ref="C2:C5" si="0">B2*2</f>' if i == 2 else '<f t="shared" si="0"/>'
            x = re.sub(rf'(<row r="{i}"[^>]*>.*?)(</row>)', rf'\1<c r="C{i}">{fx}<v>{v}</v></c>\2', x, count=1, flags=re.S)
        return x
    V.patch_sheet(p, g)


@V.fx("na_error_results")
def _(p):  # formula results #N/A (e.g. a VLOOKUP miss) and #DIV/0!, and ""
    wb, ws = V.base_book(("name", "amount", "lookup"), rows=[("Alice", 10, None), ("Bob", 20, None), ("Cy", 30, None), ("Dee", 40, None)])
    for i in range(2, 6):
        ws[f"C{i}"] = f"=B{i}"
    wb.save(p)
    res = {2: '<c r="C2" t="e"><f>B2</f><v>#N/A</v></c>', 3: '<c r="C3" t="e"><f>B3</f><v>#DIV/0!</v></c>',
           4: '<c r="C4" t="str"><f>B4</f><v></v></c>', 5: '<c r="C5"><f>B5</f><v>40</v></c>'}

    def g(x):
        for i, c in res.items():
            x = re.sub(rf'<c r="C{i}"[^>]*>.*?</c>', c, x, count=1, flags=re.S)
        return x
    V.patch_sheet(p, g)


@V.fx("entity_bomb")
def _(p):
    wb, ws = V.base_book(); wb.save(p)
    bomb = ('<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">' + "".join(f'<!ENTITY {chr(98+i)} "&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};&{chr(97+i)};">' for i in range(12)) + "]>")

    def g(x):
        x = x.replace("?>", "?>" + bomb, 1)
        return x.replace("<sheetData>", "<sheetData><!-- &m; -->", 1) if False else x.replace("<dimension", "<extLst>&m;</extLst><dimension", 1) if False else x
    V.patch_sheet(p, lambda x: re.sub(r"(<\?xml[^>]*\?>)", r"\1" + bomb.replace("\\", "\\\\"), x, count=1) if x.startswith("<?xml") else bomb + x)


def case_timing():
    import os
    import tempfile
    import openpyxl
    import pandas as pd
    from skills.excel import workbook as W
    d = tempfile.mkdtemp(); p = os.path.join(d, "big.xlsx")
    wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet("S")
    ws.append(["a", "b", "c", "d", "e"])
    for i in range(100_000):
        ws.append([i, f"n{i}", i * 2, "x", i % 7])
    wb.save(p)
    print("\n## timing, 100000 rows x 5 cols (write-only file, no <dimension>)")
    for label, fn in (("pandas.read_excel", lambda: pd.read_excel(p)), ("plain_table_problem", lambda: W.plain_table_problem(p)),
                      ("formula_cell_counts", lambda: W.formula_cell_counts(p)), ("sheet_names", lambda: W.sheet_names(p))):
        t = time.time(); r = fn(); print(f"  {label:22} {time.time() - t:6.1f}s", "" if not isinstance(r, (str, type(None), W.FormulaCounts)) else r)


if __name__ == "__main__":
    cases = {
        "measure_fails": case_measure_fails,
        "skill_doesnt_refuse": case_skill_doesnt_refuse,
        "shared_formulas": lambda: [show(f"shared formulas ({r})", V.run_one("shared_formulas", r, "template", "s1")) for r in ("dedupe", "empty")],
        "na_error_results": lambda: [show(f"#N/A results ({r})", V.run_one("na_error_results", r, "template", "s1")) for r in ("dedupe", "empty")],
        "entity_bomb": lambda: show("entity bomb in sheet XML", V.run_one("entity_bomb", "dedupe", "template", "s1")),
        "timing": case_timing,
    }
    for n, fn in cases.items():
        if not ONLY or n in ONLY:
            fn()
