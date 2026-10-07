# ruff: noqa: E501 — reviewer tooling, not product code
"""Part B check, independent of ``scripts/check_fixture_parity.py``: do the new fixture helpers make
the same workbooks as the old ones?

    python .../fixture_parity.py build TREE OUT      # TREE's tests/workbook_fixtures.py; set OPENPYXL_LXML=False for no-lxml writing
    python .../fixture_parity.py compare A B

``build`` makes the fixtures whose helpers changed (everything that goes through
``save_formula_values``, and the empty-string craft: the new helper on a tree that has it, else the
old test's inline ``str.replace``). ``compare`` checks, per part, the XML tree with text and tails
NOT stripped and attributes compared exactly (stricter than the implementer's script), openpyxl's
cell values, and this checkout's ``formula_cell_counts``, ``plain_table_problem`` and
``sheet_names`` on both.
"""

from __future__ import annotations

import os
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

HERE = Path(__file__).resolve()
REPO = HERE.parents[3]


def build(tree: Path, out: Path) -> None:
    sys.path.insert(0, str(tree))
    import openpyxl

    from tests import workbook_fixtures as fx

    out.mkdir(parents=True, exist_ok=True)
    print(f"tree={tree} openpyxl.LXML={openpyxl.LXML}", file=sys.stderr)
    fx.formatted_formula_sheet(out / "formulas_saved.xlsx", saved=True)
    fx.formatted_formula_sheet(out / "formulas_unsaved.xlsx", saved=False)
    fx.error_result_sheet(out / "error_results.xlsx")
    unsaved = fx.formatted_formula_sheet(out / "u.xlsx", saved=False)
    partly = fx.save_formula_values(unsaved, out / "partly_saved.xlsx", {"C2": 20, "C6": 60})
    if hasattr(fx, "save_empty_string_result"):
        fx.save_empty_string_result(partly, out / "empty_string_result.xlsx", "C3")
    else:  # the old test's own craft (base tests/skills/excel/test_plain_table.py)
        with zipfile.ZipFile(partly) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode()
        crafted = xml.replace('<c r="C3"><f>B3*2</f><v />', '<c r="C3" t="str"><f>B3*2</f><v></v>')
        assert crafted != xml, "the old craft didn't apply (lxml writing?)"
        fx.rewrite(partly, out / "empty_string_result.xlsx", {"xl/worksheets/sheet1.xml": crafted.encode()})
    (out / "u.xlsx").unlink()


def strict(data: bytes):
    def walk(e: ElementTree.Element):
        if e.tag.endswith(("}created", "}modified")):
            return None
        return (e.tag, tuple(sorted(e.attrib.items())), e.text, e.tail, tuple(walk(c) for c in e))

    return walk(ElementTree.fromstring(data))


def product_view(path: Path):
    sys.path.insert(0, str(REPO))
    import openpyxl

    from skills.excel import workbook

    book = openpyxl.load_workbook(path)
    cells = {s.title: [[c.value for c in r] for r in s.iter_rows()] for s in book.worksheets}
    values = openpyxl.load_workbook(path, data_only=True)
    saved = {s.title: [[c.value for c in r] for r in s.iter_rows()] for s in values.worksheets}
    return (cells, saved, workbook.formula_cell_counts(str(path)), workbook.plain_table_problem(str(path)), workbook.sheet_names(str(path)))


def compare(a: Path, b: Path) -> int:
    problems = 0
    for name in sorted(p.name for p in a.iterdir()):
        pa, pb = a / name, b / name
        with zipfile.ZipFile(pa) as za, zipfile.ZipFile(pb) as zb:
            parts = sorted(za.namelist())
            assert parts == sorted(zb.namelist()), name
            diff = [p for p in parts if strict(za.read(p)) != strict(zb.read(p))]
        same_product = product_view(pa) == product_view(pb)
        ok = not diff and same_product
        problems += not ok
        print(f"{'same' if ok else 'DIFFERENT':9} {name}: parts differing (strict) {diff}; product reads equal={same_product}")
    print(f"{problems} different")
    return 1 if problems else 0


if __name__ == "__main__":
    if sys.argv[1] == "build":
        build(Path(sys.argv[2]).resolve(), Path(sys.argv[3]))
    else:
        sys.exit(compare(Path(sys.argv[2]), Path(sys.argv[3])))
    _ = os
